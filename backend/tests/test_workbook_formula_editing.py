import io
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy.engine import Engine

from app.domain.errors import ApplicationError
from app.infrastructure.workbook_cells import WorkbookCells
from app.infrastructure.workbook_formula_editing import authored_formula, invalidate_formula_chain
from app.main import app
from tests.test_sources_api import create_project, make_client
from tests.test_workbook_editor import require_calc, sample


def test_authored_formulas_normalize_russian_functions_and_leave_quoted_text_literal() -> None:
    assert authored_formula("=СУММ(1,23;2)", set()) == "=SUM(1.23,2)"
    assert authored_formula('=IF(1>0;1,25;"1,23;a")', set()) == '=IF(1>0,1.25,"1,23;a")'
    assert authored_formula("=SUM(1,2)", set()) == "=SUM(1,2)"
    assert authored_formula('=ЕСЛИ(B2>0;СУММ(B2;5);"a;b")', set()) == '=IF(B2>0,SUM(B2,5),"a;b")'
    for formula in [
        '=WEBSERVICE("https://example.invalid")',
        '=DDE("cmd","x","y")',
        "='[9]Remote'!A1",
        "=",
        "={1,2}",
    ]:
        with pytest.raises(ApplicationError):
            authored_formula(formula, set())
    assert authored_formula("='[1]Saved'!A1+1", {(1, "Saved")}) == "='[1]Saved'!A1+1"


def test_formula_patch_adds_and_removes_formula_without_turning_literal_equals_into_code() -> None:
    package = WorkbookCells(sample())
    changed = package.patched(
        {
            package.sheets["Sales"].path: {
                "B2": ("formula", "=SUM(1,2)", False),
                "C2": ("inlineStr", "=literal", False),
            }
        }
    )
    book = load_workbook(io.BytesIO(changed), data_only=False)
    assert book["Sales"]["B2"].data_type == "f"
    assert book["Sales"]["B2"].value == "=SUM(1,2)"
    assert book["Sales"]["C2"].data_type == "s"
    assert book["Sales"]["C2"].value == "=literal"


def test_invalidating_chain_removes_references_but_preserves_original_precision_and_assets() -> (
    None
):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(sample())) as original, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for entry in original.infolist():
            value = original.read(entry.filename)
            if entry.filename == "xl/_rels/workbook.xml.rels":
                value = value.replace(
                    b"</Relationships>",
                    b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/'
                    b'2006/relationships/calcChain" Target="calcChain.xml" '
                    b'Id="rId99"/></Relationships>',
                )
            if entry.filename == "[Content_Types].xml":
                value = value.replace(
                    b"</Types>",
                    b'<Override ContentType="application/vnd.openxmlformats-officedocument.'
                    b'spreadsheetml.calcChain+xml" PartName="/xl/calcChain.xml"/></Types>',
                )
            if entry.filename == "xl/workbook.xml":
                value = value.replace(b"<calcPr ", b'<calcPr fullPrecision="0" ')
            archive.writestr(entry, value)
        archive.writestr(
            "xl/calcChain.xml",
            '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<c r="C2" i="1"/></calcChain>',
        )
    before = output.getvalue()
    result = invalidate_formula_chain(before)
    assert WorkbookCells(result).precision_as_displayed
    with ZipFile(io.BytesIO(before)) as old, ZipFile(io.BytesIO(result)) as new:
        assert "xl/calcChain.xml" not in new.namelist()
        assert b"calcChain" not in new.read("xl/_rels/workbook.xml.rels")
        assert b"calcChain" not in new.read("[Content_Types].xml")
        assert b'forceFullCalc="1"' in new.read("xl/workbook.xml")
        for name in old.namelist():
            if name not in {
                "xl/calcChain.xml",
                "xl/_rels/workbook.xml.rels",
                "[Content_Types].xml",
                "xl/workbook.xml",
            }:
                assert old.read(name) == new.read(name)


def test_array_formula_ranges_stay_structurally_protected() -> None:
    package = WorkbookCells(sample())
    original = package.content
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original)) as incoming, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for entry in incoming.infolist():
            data = incoming.read(entry.filename)
            if entry.filename == package.sheets["Sales"].path:
                data = data.replace(b"<f>B2*2</f>", b'<f t="array" ref="B2:C2">B2*2</f>')
            archive.writestr(entry, data)
    sheet = WorkbookCells(output.getvalue()).sheets["Sales"]
    assert sheet.structural_readonly("B2") and sheet.structural_readonly("C2")
    assert sheet.readonly("B2")


def shared_sample() -> bytes:
    book = load_workbook(io.BytesIO(sample()))
    book["Sales"]["B3"] = 3
    book["Sales"]["C3"] = "=$B3*2"
    book["Sales"]["C2"] = "=$B2*2"
    buffer = io.BytesIO()
    book.save(buffer)
    package = WorkbookCells(buffer.getvalue())
    content = package.patched(
        {
            package.sheets["Sales"].path: {"C2": ("n", "4", True), "C3": ("n", "6", True)},
            package.sheets["Notes"].path: {"D2": ("n", "5", True)},
        }
    )
    output = io.BytesIO()
    with ZipFile(io.BytesIO(content)) as incoming, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for entry in incoming.infolist():
            data = incoming.read(entry.filename)
            if entry.filename == package.sheets["Sales"].path:
                data = data.replace(b"<f>$B2*2</f>", b'<f t="shared" si="0" ref="C2:C3">$B2*2</f>')
                data = data.replace(b"<f>$B3*2</f>", b'<f t="shared" si="0"/>')
            archive.writestr(entry, data)
    return output.getvalue()


@pytest.mark.parametrize(
    "target,formula,expected", [(1, "=B2*10", ["20", "6"]), (2, "=B3*10", ["4", "30"])]
)
def test_editing_shared_master_or_follower_preserves_other_translated_formulas(
    db_engine: Engine, tmp_path, target, formula, expected
) -> None:
    require_calc()
    content = shared_sample()
    sheet = WorkbookCells(content).sheets["Sales"]
    assert sheet.shared_formulas == {"C2": ("0", "=$B2*2"), "C3": ("0", "=$B3*2")}
    assert not sheet.structural_readonly("C2") and not sheet.structural_readonly("C3")
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel", files={"file": ("shared.xlsx", content)}
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        page = client.get(root + "/editor", params={"resource": "Sales"}).json()
        assert page["rows"][2]["formulas"]["2"] == "=$B3*2"
        assert 2 in page["rows"][2]["editable_in_workbook_mode"]
        saved = client.post(
            root + "/editor",
            json={
                "expected_revision": update["revision"],
                "request_id": str(uuid.uuid4()),
                "resource_id": "Sales",
                "cells": [
                    {
                        "row_id": f"r:{target}",
                        "column": 2,
                        "value": formula,
                        "formula": True,
                        "override_protection": True,
                    }
                ],
            },
        )
        assert saved.status_code == 200, saved.text
        prepared = client.post(
            root + "/editor/prepare",
            json={"expected_revision": saved.json()["revision"], "request_id": str(uuid.uuid4())},
        )
        assert prepared.status_code == 200, prepared.text
        staged = client.get(root + "/editor", params={"resource": "Sales"}).json()
        assert [row["values"][2] for row in staged["rows"][1:]] == expected
        other = 2 if target == 1 else 1
        assert staged["rows"][other]["formulas"]["2"] == f"=$B{other + 1}*2"
        original = client.get(f"/api/v1/sources/{source}/sheets/Sales/preview").json()
        assert [row["col_3"] for row in original["rows"]] == [4, 6]
    finally:
        client.close()
        app.dependency_overrides.clear()


def test_explicit_formula_and_protected_value_edits_recalculate_before_atomic_apply(
    db_engine: Engine, tmp_path
) -> None:
    require_calc()
    client = make_client(db_engine, tmp_path)
    try:
        book = load_workbook(io.BytesIO(sample()))
        book["Sales"].protection.sheet = True
        book["Notes"].protection.sheet = True
        stream = io.BytesIO()
        book.save(stream)
        # Restore cached results after fixture serialization.
        package = WorkbookCells(stream.getvalue())
        content = package.patched(
            {
                package.sheets["Sales"].path: {"C2": ("n", "4", True)},
                package.sheets["Notes"].path: {"D2": ("n", "5", True)},
            }
        )
        project = str(create_project(db_engine))
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel", files={"file": ("protected.xlsx", content)}
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        page = client.get(root + "/editor", params={"resource": "Sales"}).json()
        assert page["rows"][1]["editable_in_workbook_mode"] == [0, 1, 2]

        def save(resource, revision, cells):
            return client.post(
                root + "/editor",
                json={
                    "expected_revision": revision,
                    "request_id": str(uuid.uuid4()),
                    "resource_id": resource,
                    "cells": cells,
                },
            )

        assert (
            save(
                "Sales",
                update["revision"],
                [{"row_id": "r:1", "column": 2, "value": "=B2+5", "formula": True}],
            ).status_code
            == 422
        )
        blocked = save(
            "Sales",
            update["revision"],
            [
                {
                    "row_id": "r:1",
                    "column": 2,
                    "value": '=WEBSERVICE("https://example.invalid")',
                    "formula": True,
                    "override_protection": True,
                }
            ],
        )
        assert blocked.status_code == 422
        assert client.get(root + "/editor").json()["revision"] == update["revision"]
        saved = save(
            "Sales",
            update["revision"],
            [
                {"row_id": "r:1", "column": 1, "value": "3", "override_protection": True},
                {
                    "row_id": "r:1",
                    "column": 2,
                    "value": "=СУММ(B2;5)",
                    "formula": True,
                    "override_protection": True,
                },
            ],
        )
        assert saved.status_code == 200, saved.text
        assert saved.json()["rows"][1]["formulas"]["2"] == "=SUM(B2,5)"
        assert saved.json()["rows"][1]["values"][2] == ""
        assert saved.json()["rows"][1]["pending_formulas"] == [2]
        # Writing into a merged follower must remain an atomic rejection, even in this mode.
        follower = save(
            "Notes",
            saved.json()["revision"],
            [{"row_id": "r:0", "column": 1, "value": "bad", "override_protection": True}],
        )
        assert follower.status_code == 422
        notes = save(
            "Notes",
            saved.json()["revision"],
            [
                {
                    "row_id": "r:0",
                    "column": 0,
                    "value": "New merged heading",
                    "override_protection": True,
                },
                {
                    "row_id": "r:1",
                    "column": 2,
                    "value": "=Sales!B2+10",
                    "formula": True,
                    "override_protection": True,
                },
            ],
        )
        assert notes.status_code == 200, notes.text
        assert (
            client.get(f"/api/v1/sources/{source}/sheets/Sales/preview").json()["rows"][0]["col_2"]
            == 2
        )
        prepared = client.post(
            root + "/editor/prepare",
            json={"expected_revision": notes.json()["revision"], "request_id": str(uuid.uuid4())},
        )
        assert prepared.status_code == 200, prepared.text
        sales = client.get(root + "/editor", params={"resource": "Sales"}).json()
        assert sales["rows"][1]["values"] == ["0001", "3", "8"]
        notes_page = client.get(root + "/editor", params={"resource": "Notes"}).json()
        assert notes_page["rows"][1]["values"][2:] == ["13", "9"]
        assert (
            client.post(
                root + "/check", json={"expected_revision": prepared.json()["revision"]}
            ).status_code
            == 200
        )
        checked = client.get(f"/api/v1/projects/{project}/updates").json()[0]
        assert checked["check"]["issues"] == []
        assert (
            client.post(
                root + "/apply", json={"expected_revision": checked["revision"]}
            ).status_code
            == 200
        )
        active = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        file = client.get(f"/api/v1/updates/{active['id']}/files/current")
        restored = load_workbook(io.BytesIO(file.content), data_only=False)
        assert restored["Sales"]["C2"].value == "=SUM(B2,5)"
        assert restored["Notes"]["C2"].value == "=Sales!B2+10"
        assert restored["Notes"]["A1"].value == "New merged heading"
        assert str(restored["Notes"].merged_cells) == "A1:C1"
        assert restored["Sales"].protection.sheet and restored["Notes"].protection.sheet
    finally:
        client.close()
        app.dependency_overrides.clear()
