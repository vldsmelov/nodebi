import io
import uuid
from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Color, Font, PatternFill, Side
from openpyxl.utils.datetime import MAC_EPOCH
from sqlalchemy.engine import Engine

from app.infrastructure.workbook_cells import WorkbookCells
from app.main import app
from tests.test_sources_api import create_project, make_client


@pytest.mark.parametrize("mac_dates", [False, True])
def test_presentation_keeps_numeric_semantics_styles_dates_codes_and_empty_cells_separate(
    db_engine: Engine, tmp_path, mac_dates: bool
) -> None:
    book = Workbook()
    if mac_dates:
        book.epoch = MAC_EPOCH
    sheet = book.active
    assert sheet is not None
    sheet.title = "Inputs"
    sheet.append(["Code", "Amount", "Formula", "Date", "Empty"])
    sheet.append(["00009007199254740993", 1234.5678, "=B2*2", datetime(2025, 3, 15), None])
    sheet["B2"].number_format = "#,##0.00"
    sheet["B2"].font = Font(name="Calibri", size=12, color=Color(theme=4), bold=True)
    sheet["B2"].fill = PatternFill("solid", fgColor=Color(indexed=6))
    sheet["B2"].border = Border(bottom=Side(style="thin", color="00112233"))
    sheet["D2"].number_format = "dd.mm.yyyy"
    stream = io.BytesIO()
    book.save(stream)
    package = WorkbookCells(stream.getvalue())
    content = package.patched({package.sheets["Inputs"].path: {"C2": ("n", "2469.1356", True)}})
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel",
            files={"file": ("presentation.xlsx", content)},
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        workspace = client.get(root + "/editor/workspace").json()
        page = workspace["page"]
        assert page["workbook"] and page["date_1904"] == mac_dates
        row = page["rows"][1]
        assert row["values"] == [
            "00009007199254740993",
            "1234.5678",
            "2469.1356",
            "2025-03-15T00:00:00",
            "",
        ]
        assert set(row["number_values"]) == {"1", "2", "3"}
        assert row["number_values"]["1"] == "1234.5678"
        assert row["number_values"]["2"] == "2469.1356"
        assert row["formulas"]["2"] == "=B2*2"
        style = workspace["styles"][row["style_ids"]["1"]]
        assert style["bold"] and style["font"] == "Calibri"
        assert style["number_format"] == "#,##0.00" and style["alignment"] == "general"
        assert style["color"] == "#4F81BD"
        assert style["background"] == "#FF00FF"
        assert style["borders"] == {"bottom": {"line": "thin", "color": "#112233"}}
        assert client.get(root + "/files/current").content == content
    finally:
        client.close()
        app.dependency_overrides.clear()
