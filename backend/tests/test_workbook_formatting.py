import io
import uuid
from copy import deepcopy
from pathlib import Path
from xml.etree.ElementTree import SubElement, tostring
from zipfile import ZipFile

import pytest
from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, Side
from sqlalchemy.engine import Engine

from app.domain.ingestion import ImportField, ImportProfile
from app.domain.workbook_formats import CellFormatPatch
from app.infrastructure.workbook_cells import NS, WorkbookCells
from app.infrastructure.workbook_formatting import format_book
from app.infrastructure.worksheet_input_precision import preserve_numbers
from app.main import app
from tests.test_dataset_updates import check, import_data
from tests.test_sources_api import create_project, make_client


def sample(formula=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Inputs"
    sheet.append(["Code", "Amount", "Result"])
    sheet.append(["0001", 10.25, "=B2*2" if formula else "text"])
    for coordinate in ("B2", "C2"):
        sheet[coordinate].font = Font(name="Calibri", size=11, color="FF123456")
        sheet[coordinate].border = Border(bottom=Side(style="thin", color="FF654321"))
        sheet[coordinate].number_format = "#,##0.00"
    stream = io.BytesIO()
    book.save(stream)
    package = WorkbookCells(stream.getvalue())
    return (
        package.patched({package.sheets["Inputs"].path: {"C2": ("n", "20.5", True)}})
        if formula
        else stream.getvalue()
    )


def test_appended_formats_preserve_original_assets_records_values_formulas_and_precision():
    content = sample(True)
    patch = CellFormatPatch(
        bold=True,
        italic=True,
        color="#008800",
        background="#FFFF00",
        alignment="right",
        wrap=True,
        size=16,
    )
    result = format_book(
        content, {"Inputs": {"B2": patch, "C2": patch, "A2": CellFormatPatch(background="none")}}
    )
    before, after = WorkbookCells(content), WorkbookCells(result)
    sheet = after.sheets["Inputs"]
    assert sheet.style_ids["B2"] == sheet.style_ids["C2"]  # Same source style/patch reuses one xf.
    assert before.precision_as_displayed == after.precision_as_displayed
    for coordinate in before.sheets["Inputs"].cells:
        old = deepcopy(before.sheets["Inputs"].cells[coordinate])
        new = deepcopy(sheet.cells[coordinate])
        old.attrib.pop("s", None)
        new.attrib.pop("s", None)
        assert tostring(old) == tostring(new)
    with ZipFile(io.BytesIO(content)) as a, ZipFile(io.BytesIO(result)) as b:
        for name in a.namelist():
            if name not in {"xl/styles.xml", sheet.path}:
                assert a.read(name) == b.read(name)
        first, second = fromstring(a.read("xl/styles.xml")), fromstring(b.read("xl/styles.xml"))
        for container in ("fonts", "fills", "cellXfs"):
            for original, copied in zip(
                first.find(NS + container), second.find(NS + container), strict=False
            ):
                assert tostring(original) == tostring(copied)
    book = load_workbook(io.BytesIO(result), data_only=False)
    assert book.active["C2"].value == "=B2*2"
    cell = book.active["B2"]
    assert cell.value == 10.25 and cell.font.bold and cell.font.italic and cell.font.sz == 16
    assert cell.font.name == "Calibri" and cell.font.color.rgb == "FF008800"
    assert cell.fill.fgColor.rgb == "FFFFFF00" and cell.alignment.wrap_text
    assert cell.number_format == "#,##0.00" and cell.border.bottom.color.rgb == "FF654321"


@pytest.mark.parametrize("physical", [False, True])
def test_format_only_drafts_reopen_prepare_and_apply_without_changing_values(
    db_engine: Engine, tmp_path: Path, monkeypatch, physical
):
    monkeypatch.setattr("app.services.workbook_editor.recalculate", lambda content: content)
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        original = sample(physical)
        package = WorkbookCells(original)
        original = package.patched(
            {package.sheets["Inputs"].path: {"A2": ("n", "9007199254740993", False)}}
        )
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel",
            files={"file": ("format.xlsx", original)},
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        workspace = client.get(root + "/editor/workspace").json()
        assert workspace["page"]["can_format"] and workspace["page"]["workbook"] == physical
        body = {
            "resource_id": "Inputs",
            "expected_revision": update["revision"],
            "request_id": str(uuid.uuid4()),
            "formats": [
                {
                    "row_id": "r:1" if physical else "r:0",
                    "column": 1,
                    "patch": {"bold": True, "background": "#FFFF00"},
                }
            ],
        }
        response = client.post(root + "/editor", json=body)
        assert response.status_code == 200, response.text
        saved = response.json()
        assert client.post(root + "/editor", json=body).json() == saved
        row = next(
            r
            for r in client.get(root + "/editor").json()["rows"]
            if r["id"] == body["formats"][0]["row_id"]
        )
        assert row["formats"]["1"]["bold"] and row["values"][1] == "10.25"
        assert client.get(root + "/files/current").content == original
        prepared = client.post(
            root + "/editor/prepare",
            json={"expected_revision": saved["revision"], "request_id": str(uuid.uuid4())},
        )
        assert prepared.status_code == 200, prepared.text
        checked = check(client, project, prepared.json())
        assert not any(i["severity"] == "error" for i in checked["check"]["issues"])
        applied = client.post(
            root + "/apply",
            json={"expected_revision": checked["revision"], "acknowledge_warnings": True},
        )
        assert applied.status_code == 200, applied.text
        next_update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        content = client.get(f"/api/v1/updates/{next_update['id']}/files/current").content
        assert (
            WorkbookCells(content).sheets["Inputs"].cells["A2"].findtext(NS + "v")
            == "9007199254740993"
        )
        cell = load_workbook(io.BytesIO(content)).active["B2"]
        assert cell.value == 10.25 and cell.font.bold and cell.fill.fgColor.rgb == "FFFFFF00"
    finally:
        client.close()
        app.dependency_overrides.clear()


def test_csv_rejects_formatting_without_leaving_a_partial_draft(db_engine, tmp_path):
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        source, _ = import_data(client, project, b"code,amount\n0001,10\n")
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source["id"], "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        assert not client.get(root + "/editor/workspace").json()["page"]["can_format"]
        response = client.post(
            root + "/editor",
            json={
                "expected_revision": update["revision"],
                "resource_id": "table",
                "request_id": str(uuid.uuid4()),
                "cells": [{"row_id": "r:0", "column": 1, "value": "99"}],
                "formats": [{"row_id": "r:0", "column": 1, "patch": {"bold": True}}],
            },
        )
        assert response.status_code == 422, response.text
        assert client.get(root + "/editor").json()["rows"][0]["values"][1] == "10"
    finally:
        client.close()
        app.dependency_overrides.clear()


@pytest.mark.parametrize("spacing", [False, True])
def test_new_empty_formatted_cell_and_prefixed_styles_remain_valid(spacing):
    import re

    content = sample()
    book = load_workbook(io.BytesIO(content))
    book.active["A3"] = "Tail"
    buffer = io.BytesIO()
    book.save(buffer)
    prefixed = io.BytesIO()
    with ZipFile(buffer) as source, ZipFile(prefixed, "w") as target:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                xml = value.decode().replace("xmlns=", "xmlns:s=")
                if spacing:
                    xml = xml.replace(' count="', ' count = "')
                value = re.sub(r"<(/?)([A-Za-z_]\w*)", r"<\1s:\2", xml).encode()
            elif spacing and item.filename.startswith("xl/worksheets/"):
                value = value.replace(b' r="', b' r = "').replace(b' s="', b' s = "')
            target.writestr(item, value)
    result = format_book(
        prefixed.getvalue(),
        {"Inputs": {"C3": CellFormatPatch(bold=True), "B2": CellFormatPatch(italic=True)}},
    )
    sheet = load_workbook(io.BytesIO(result)).active
    assert sheet["C3"].value in {None, ""} and sheet["C3"].font.bold
    assert sheet["A3"].value == "Tail" and sheet["B2"].value == 10.25 and sheet["B2"].font.italic


def test_explicit_false_font_flags_are_not_shown_as_enabled():
    buffer = io.BytesIO()
    with ZipFile(io.BytesIO(sample())) as source, ZipFile(buffer, "w") as target:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                root = fromstring(value)
                font = root.find(NS + "fonts/" + NS + "font")
                SubElement(font, NS + "b", {"val": "0"})
                SubElement(font, NS + "i", {"val": "false"})
                value = tostring(root)
            target.writestr(item, value)
    style = WorkbookCells(buffer.getvalue()).display_styles[0]
    assert style["bold"] is False and style["italic"] is False


@pytest.mark.parametrize("move", [False, True])
def test_numeric_precision_is_kept_for_other_cells_moved_rows_and_explicit_integer_edits(move):
    book = load_workbook(io.BytesIO(sample()))
    book.active.append([2, 20, "Tail"])
    book.create_sheet("Unused")["A1"] = 3
    buffer = io.BytesIO()
    book.save(buffer)
    package = WorkbookCells(buffer.getvalue())
    content = package.patched(
        {
            package.sheets["Inputs"].path: {
                "A2": ("n", "9007199254740993", False),
                "A3": ("n", "9999999999999999", False),
            },
            package.sheets["Unused"].path: {"A1": ("n", "12345678901234567", False)},
        }
    )
    package = WorkbookCells(content)
    rendered = load_workbook(io.BytesIO(content))
    if move:
        rendered["Inputs"].delete_rows(2)
        rendered["Inputs"]["A2"] = 9999999999999997
        state = {"cells": {"r:1": {"0": "9999999999999997"}}, "deleted": ["r:0"], "added": []}
    else:
        rendered["Inputs"]["B2"] = 11.75
        state = {"cells": {"r:0": {"1": "11.75"}}, "deleted": [], "added": []}
    output = io.BytesIO()
    rendered.save(output)
    profile = ImportProfile(
        resource_id="Inputs",
        name="Inputs",
        fields=[
            ImportField(field_id="code", source_field="col_1", name="Code", logical_type="integer")
        ],
    )
    result = preserve_numbers(
        package,
        output.getvalue(),
        [profile],
        {"Inputs": state},
        {"Inputs": 2},
        lambda p, c, v: int(v) if c == 0 else float(v),
    )
    changed = WorkbookCells(result)
    assert changed.sheets["Inputs"].cells["A2"].findtext(NS + "v") == (
        "9999999999999997" if move else "9007199254740993"
    )
    assert changed.sheets["Unused"].cells["A1"].findtext(NS + "v") == "12345678901234567"
    if not move:
        assert changed.sheets["Inputs"].cells["A3"].findtext(NS + "v") == "9999999999999999"
        assert changed.sheets["Inputs"].cells["B2"].findtext(NS + "v") == "11.75"


def test_protected_formatting_requires_explicit_override_and_preserves_protection(
    db_engine, tmp_path, monkeypatch
):
    monkeypatch.setattr("app.services.workbook_editor.recalculate", lambda content: content)
    book = load_workbook(io.BytesIO(sample(True)))
    book.active.protection.sheet = True
    buffer = io.BytesIO()
    book.save(buffer)
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel",
            files={"file": ("locked.xlsx", buffer.getvalue())},
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        request = {
            "expected_revision": update["revision"],
            "request_id": str(uuid.uuid4()),
            "resource_id": "Inputs",
            "formats": [{"row_id": "r:1", "column": 1, "patch": {"bold": True}}],
        }
        rejected = client.post(root + "/editor", json=request)
        assert (
            rejected.status_code == 422
            and rejected.json()["detail"]["code"] == "EDITOR_CELL_READONLY"
        )
        request["formats"][0]["override_protection"] = True
        accepted = client.post(root + "/editor", json=request)
        assert accepted.status_code == 200, accepted.text
        prepared = client.post(
            root + "/editor/prepare",
            json={
                "expected_revision": accepted.json()["revision"],
                "request_id": str(uuid.uuid4()),
            },
        )
        assert prepared.status_code == 200, prepared.text
        assert client.get(root + "/files/current").content == buffer.getvalue()
        assert client.get(root + "/editor/workspace").json()["page"]["rows"][1]["format_readonly"]
    finally:
        client.close()
        app.dependency_overrides.clear()
