import uuid
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy.engine import Engine

from app.domain.ingestion import ReadOptions
from app.domain.workbooks import WorksheetLayout, WorksheetLoadOptions
from app.infrastructure.connectors import FileAdapter
from app.infrastructure.excel import ExcelWorkbookError, OpenPyxlWorkbookReader
from app.infrastructure.storage import LocalStorageProvider
from tests.test_dataset_updates import build_report, check, stage_update
from tests.test_sources_api import create_project, make_client


def workbook_file(storage: LocalStorageProvider, merged: bool = False) -> str:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Outside", "Report", None, "Outside"])
    sheet.append(["Ignored", "Code", "Amount", "Ignored"])
    sheet.append(["=1+1", "001", "0.10", "=1+1"])
    sheet.append(["Ignored", "002", "0.20", "Ignored"])
    sheet.append(["Ignored", "Total", "0.30", "Ignored"])
    if merged:
        sheet.merge_cells("B1:C1")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return storage.save(stream, suffix=".xlsx")


def test_range_preserves_physical_identity_and_ignores_outside_formulas(tmp_path: Path) -> None:
    storage = LocalStorageProvider(tmp_path)
    key = workbook_file(storage)
    options = ReadOptions(cell_range=" b2:c5 ", drop_tail_rows=1)
    assert options.cell_range == "B2:C5"
    data = list(FileAdapter("excel", storage).read_batches(key, "Sheet", options))[0]
    assert [c.column_id for c in data.schema.columns] == ["col_2", "col_3"]
    assert [c.metadata["source_position"] for c in data.schema.columns] == [1, 2]
    assert data.frame.rows() == [("001", "0.10"), ("002", "0.20")]
    reader = OpenPyxlWorkbookReader(storage)
    preview = reader.preview_sheet(
        key,
        "Sheet",
        WorksheetLoadOptions(layout=WorksheetLayout(cell_range="B2:C5", drop_tail_rows=1)),
    )
    assert preview.rows == [{"col_2": "001", "col_3": "0.10"}, {"col_2": "002", "col_3": "0.20"}]
    with pytest.raises(ExcelWorkbookError, match="Formula has no saved result"):
        reader.load_sheet(key, "Sheet")
    with pytest.raises(ValueError, match="only to Excel"):
        list(FileAdapter("csv", storage).read_batches(key, "table", options))


def test_range_multi_headers_empty_data_and_bounds(tmp_path: Path) -> None:
    storage = LocalStorageProvider(tmp_path)
    key = workbook_file(storage, merged=True)
    reader = OpenPyxlWorkbookReader(storage)
    data = reader.load_sheet(key, "Sheet", WorksheetLayout(cell_range="B1:C4", header_rows=2))
    assert [c.display_name for c in data.schema.columns] == ["Report / Code", "Report / Amount"]
    assert data.row_count == 2
    assert reader.load_sheet(key, "Sheet", WorksheetLayout(cell_range="B2:C2")).row_count == 0
    with pytest.raises(ExcelWorkbookError, match="merged cell"):
        reader.load_sheet(key, "Sheet", WorksheetLayout(cell_range="C1:C4"))
    with pytest.raises(ExcelWorkbookError, match="beyond"):
        reader.load_sheet(key, "Sheet", WorksheetLayout(cell_range="B2:C6"))


@pytest.mark.parametrize(
    "value", ["", "B2", "Sheet!B2:C4", "A:A", "B4:A2", "A0:B2", "A1:B250001", "A1:XFE2", "A1:XFD2"]
)
def test_invalid_ranges_are_rejected(value: str) -> None:
    with pytest.raises(ValueError):
        ReadOptions(cell_range=value)


def test_header_count_must_fit_range() -> None:
    with pytest.raises(ValueError, match="Header rows must fit"):
        ReadOptions(cell_range="B2:C2", header_rows=2)


@pytest.mark.parametrize("relocated", [False, True])
def test_range_api_report_refresh_and_rollback(
    db_engine: Engine, tmp_path: Path, relocated: bool
) -> None:
    project = str(create_project(db_engine))
    client = make_client(db_engine, tmp_path)

    def content(amount: str, moved: bool = False) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["Report", "ignored", "ignored", "ignored"])
        sheet.append(["Ignored", "Code", "Amount", "Ignored"])
        sheet.append(["=1+1", "001", amount, "=1+1"])
        sheet.append(["Total", "ignored", "bad", "ignored"])
        if moved:
            sheet.delete_rows(1, 4)
            sheet["A1"] = "Report"
            sheet["D4"], sheet["E4"] = "Money", "Key"
            sheet["D5"], sheet["E5"] = "Note", "Note"
            sheet["D6"], sheet["E6"] = amount, "001"
            sheet["D7"], sheet["E7"] = "Total", "Total"
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    uploaded = client.post(
        f"/api/v1/projects/{project}/imports",
        data={"connector": "excel"},
        files={"file": ("range.xlsx", content("0.10"))},
    )
    assert uploaded.status_code == 201, uploaded.text
    source = uploaded.json()
    inspected = client.post(
        f"/api/v1/imports/{source['id']}/inspect",
        json={"resource_id": "Sheet", "options": {"cell_range": "B2:C3"}},
    )
    assert inspected.status_code == 200, inspected.text
    profile = inspected.json()["profile"]
    assert [f["source_field"] for f in profile["fields"]] == ["col_2", "col_3"]
    profile["fields"][1]["logical_type"] = "decimal"
    applied = client.post(
        f"/api/v1/imports/{source['id']}/apply",
        json={"profile": profile, "expected_revision": 1, "request_id": str(uuid.uuid4())},
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["profile"]["options"]["cell_range"] == "B2:C3"
    report = build_report(client, project, source, applied.json())
    result_url = f"/api/v1/visualizations/{report['visualization_id']}/result"
    assert client.post(result_url, json={}).json()["rows"][0]["Revenue"] == "0.10"
    assert client.post(f"/api/v1/projects/{project}/checks").status_code == 200
    job = client.get(f"/api/v1/projects/{project}/report-jobs").json()[0]
    assert job["status"] == "succeeded", job
    frozen = client.post(
        f"/api/v1/projects/{project}/releases",
        json={
            "name": "Range",
            "period": "October",
            "dashboard_ids": [report["dashboard_id"]],
            "expected_token": job["check"]["token"],
        },
    )
    assert frozen.status_code == 200, frozen.text
    release = client.get(f"/api/v1/projects/{project}/releases").json()[0]
    staged = stage_update(client, project, source, content("0.40", relocated), "next.xlsx")
    if relocated:
        binding = {
            "resource_id": "Sheet",
            "target_resource_id": "Sheet",
            "options": {"cell_range": "D4:E999", "skip_data_rows": 1, "drop_tail_rows": 1},
            "columns": [
                {"logical_id": profile["fields"][0]["field_id"], "physical_id": "col_5"},
                {"logical_id": profile["fields"][1]["field_id"], "physical_id": "col_4"},
            ],
        }
        staged["items"][0]["dataset_bindings"] = [binding]
        saved = client.put(
            f"/api/v1/updates/{staged['id']}",
            json={"expected_revision": staged["revision"], "items": staged["items"]},
        )
        assert saved.status_code == 200, saved.text
        failed = check(client, project, saved.json())
        assert any(i["severity"] == "error" for i in failed["check"]["issues"])
        refused = client.post(
            f"/api/v1/updates/{failed['id']}/apply",
            json={"expected_revision": failed["revision"], "acknowledge_warnings": True},
        )
        assert refused.status_code == 422, refused.text
        assert refused.json()["detail"]["code"] == "CHECK_FAILED"
        assert client.post(result_url, json={}).json()["rows"][0]["Revenue"] == "0.10"
        binding["options"]["cell_range"] = "D4:E7"
        failed["items"][0]["dataset_bindings"] = [binding]
        saved = client.put(
            f"/api/v1/updates/{failed['id']}",
            json={"expected_revision": failed["revision"], "items": failed["items"]},
        )
        assert saved.status_code == 200, saved.text
        staged = saved.json()
    update = check(client, project, staged)
    assert not update["check"]["issues"], update
    refreshed = client.post(
        f"/api/v1/updates/{update['id']}/apply", json={"expected_revision": update["revision"]}
    )
    assert refreshed.status_code == 200, refreshed.text
    assert client.post(result_url, json={}).json()["rows"][0]["Revenue"] == "0.40"
    copied = client.post(
        f"/api/v1/releases/{release['id']}/restore", json={"name": "Restored range"}
    )
    assert copied.status_code == 200, copied.text
    restored_project = copied.json()["id"]
    datasets = client.get(f"/api/v1/projects/{restored_project}/datasets").json()
    assert datasets[0]["profile"]["options"]["cell_range"] == "B2:C3"
    restored_visual = client.get(f"/api/v1/projects/{restored_project}/visualizations").json()[
        "items"
    ][0]
    restored_result = client.post(f"/api/v1/visualizations/{restored_visual['id']}/result", json={})
    assert restored_result.status_code == 200, restored_result.text
    assert restored_result.json()["rows"][0]["Revenue"] == "0.10"
    rollback = client.post(
        f"/api/v1/updates/{update['id']}/rollback",
        json={"expected_revision": refreshed.json()["revision"]},
    )
    assert rollback.status_code == 200, rollback.text
    checked = check(client, project, rollback.json())
    restored = client.post(
        f"/api/v1/updates/{checked['id']}/apply", json={"expected_revision": checked["revision"]}
    )
    assert restored.status_code == 200, restored.text
    assert client.post(result_url, json={}).json()["rows"][0]["Revenue"] == "0.10"
