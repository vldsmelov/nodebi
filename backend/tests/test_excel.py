from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from app.domain.workbooks import LogicalType, WorksheetLoadOptions
from app.infrastructure.excel import ExcelWorkbookError, OpenPyxlWorkbookReader
from app.infrastructure.storage import LocalStorageProvider


@pytest.mark.parametrize(
    "headers",
    [
        ["a", "a", "a__2"],
        ["a__2", "a", "a", "a__3"],
        [None, "Column 1", "Column 1__2"],
        ["a", "a", "a", "a__2", "a__3"],
        [None, None, None],
    ],
)
def test_header_suffixes_cannot_collide(tmp_path: Path, headers: list[str | None]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    sheet.append(list(range(len(headers))))
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(stream, suffix=".xlsx")
    reader = OpenPyxlWorkbookReader(storage)
    preview = reader.preview_sheet(key, sheet.title, WorksheetLoadOptions())
    dataset = reader.load_sheet(key, sheet.title)
    names = [column.engine_name for column in preview.columns]
    assert len(set(names)) == len(headers)
    assert dataset.frame.columns == names
    assert dataset.frame.row(0) == tuple(range(len(headers)))


def test_full_sheet_inference_handles_late_mixed_values_and_nulls(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["mixed", "number"])
    for number in range(101):
        sheet.append([number, number])
    sheet.append(["late text", 1.5])
    sheet.append([None, None])
    sheet.append(["last", 2])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(stream, suffix=".xlsx")
    dataset = OpenPyxlWorkbookReader(storage).load_sheet(key, sheet.title)
    assert dataset.frame["mixed"].to_list() == [*map(str, range(101)), "late text", None, "last"]
    assert dataset.schema.by_id("col_1").logical_type == LogicalType.STRING
    assert dataset.schema.by_id("col_2").logical_type == LogicalType.FLOAT
    assert dataset.frame["number"][101] == 1.5


def workbook_bytes() -> BytesIO:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["name", "amount", "amount", None, "when", "active", "mixed"])
    sales.append(["Alpha", 2, 2.5, None, date(2026, 1, 2), True, 1])
    sales.append(["Бета", None, 3.5, "value", datetime(2026, 2, 3, 4, 5), False, "x"])
    workbook.create_sheet("Empty")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return stream


def test_inspection_and_preview_cover_excel_edge_cases(tmp_path: Path) -> None:
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(workbook_bytes(), suffix=".xlsx")
    reader = OpenPyxlWorkbookReader(storage)

    info = reader.inspect(key)
    preview = reader.preview_sheet(key, "Sales", WorksheetLoadOptions(limit=1))
    empty = reader.preview_sheet(key, "Empty", WorksheetLoadOptions())

    assert [sheet.name for sheet in info.sheets] == ["Sales", "Empty"]
    assert [column.engine_name for column in preview.columns[:4]] == [
        "name",
        "amount",
        "amount__2",
        "Column 4",
    ]
    assert preview.columns[1].logical_type is LogicalType.INTEGER
    assert preview.columns[3].nullable
    assert preview.columns[4].logical_type is LogicalType.DATETIME
    assert preview.columns[5].logical_type is LogicalType.BOOLEAN
    assert preview.rows[0]["col_5"] == "2026-01-02T00:00:00"
    assert preview.truncated
    assert empty.columns == []
    assert empty.rows == []


def test_mixed_column_falls_back_to_string(tmp_path: Path) -> None:
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(workbook_bytes(), suffix=".xlsm")

    preview = OpenPyxlWorkbookReader(storage).preview_sheet(
        key, "Sales", WorksheetLoadOptions(limit=10)
    )

    assert preview.columns[-1].logical_type is LogicalType.STRING


def test_sheet_loads_as_source_independent_dataset(tmp_path: Path) -> None:
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(workbook_bytes(), suffix=".xlsx")

    dataset = OpenPyxlWorkbookReader(storage).load_sheet(key, "Sales")

    assert dataset.row_count == 2
    assert dataset.frame.columns[1:3] == ["amount", "amount__2"]
    assert dataset.schema.by_id("col_2").nullable
    assert dataset.schema.by_display_name("amount")[1].engine_name == "amount__2"
    assert dataset.metadata["sheet_name"] == "Sales"


@pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
@pytest.mark.parametrize("unsafe", [True, False])
def test_rejects_unsafe_and_malformed_xml(tmp_path: Path, part: str, unsafe: bool) -> None:
    modified = BytesIO()
    with ZipFile(workbook_bytes()) as original, ZipFile(modified, "w") as archive:
        for member in original.infolist():
            data = original.read(member.filename)
            if member.filename == part:
                data = (
                    b'<!DOCTYPE workbook [<!ENTITY example "harmless">]>' + data
                    if unsafe
                    else b"<broken>"
                )
            archive.writestr(member, data)
    modified.seek(0)
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(modified, suffix=".xlsx")
    reader = OpenPyxlWorkbookReader(storage)
    with pytest.raises(ExcelWorkbookError):
        reader.load_sheet(key, "Sales")
    with pytest.raises(ExcelWorkbookError):
        reader.preview_sheet(key, "Sales", WorksheetLoadOptions())


def test_preview_rejects_excessive_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["field"] * 1025)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    storage = LocalStorageProvider(tmp_path)
    key = storage.save(stream, suffix=".xlsx")
    with pytest.raises(ExcelWorkbookError, match="1024-column"):
        OpenPyxlWorkbookReader(storage).preview_sheet(key, "Sheet", WorksheetLoadOptions())
