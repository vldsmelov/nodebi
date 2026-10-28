import io
import uuid
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from sqlalchemy.engine import Engine

from app.domain.errors import ApplicationError
from app.infrastructure.spreadsheet_recalculation import office_binary, recalculate, safe_formula
from app.infrastructure.workbook_cells import WorkbookCells
from app.main import app
from tests.test_sources_api import create_project, make_client


def sample() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Sales"
    sheet.append(["Code", "Quantity", "Total"])
    sheet.append(["0001", 2, "=B2*2"])
    notes = book.create_sheet("Notes")
    notes["A1"] = "Keep this layout"
    notes.merge_cells("A1:C1")
    notes["D2"] = "=Sales!C2+1"
    stream = io.BytesIO()
    book.save(stream)
    package = WorkbookCells(stream.getvalue())
    return package.patched(
        {
            package.sheets["Sales"].path: {"C2": ("n", "4", True)},
            package.sheets["Notes"].path: {"D2": ("n", "5", True)},
        }
    )


def test_surgical_patch_keeps_original_parts_formulas_merges_and_literal_strings() -> None:
    content = sample()
    package = WorkbookCells(content)
    result = package.patched(
        {
            package.sheets["Sales"].path: {
                "A2": ("inlineStr", '=WEBSERVICE("https://example.invalid")', False),
                "B2": ("n", "7", False),
                "C2": ("n", "14", True),
            },
            package.sheets["Notes"].path: {"B2": ("inlineStr", "001", False)},
        }
    )
    book = load_workbook(io.BytesIO(result), data_only=False)
    assert book["Sales"]["A2"].data_type == "s"
    assert book["Sales"]["B2"].value == 7
    assert book["Sales"]["C2"].value == "=B2*2"
    assert str(book["Notes"].merged_cells) == "A1:C1"
    assert book["Notes"]["B2"].value == "001"
    with ZipFile(io.BytesIO(content)) as old, ZipFile(io.BytesIO(result)) as new:
        for name in old.namelist():
            if not name.startswith("xl/worksheets/"):
                assert old.read(name) == new.read(name)


def test_formatting_tails_and_unlocked_inputs_on_protected_sheets() -> None:
    book = load_workbook(io.BytesIO(sample()))
    sheet = book["Sales"]
    sheet.protection.sheet = True
    sheet["A2"].protection = Protection(locked=False)
    sheet["C2"].protection = Protection(locked=False)
    sheet["XFD1048576"].number_format = "0.00"
    output = io.BytesIO()
    book.save(output)
    package = WorkbookCells(output.getvalue())
    target = package.sheets["Sales"]
    assert (target.rows, target.columns) == (2, 3)
    assert not target.readonly("A2")
    assert target.readonly("B2")
    assert target.readonly("C2")  # Formula remains protected even if Excel unlocked it.


@pytest.mark.parametrize(
    "formula", ['=WEBSERVICE("https://example.invalid")', '=DDE("cmd","x","y")', "='[1]Sheet'!A1"]
)
def test_unsafe_or_uncached_external_formulas_are_blocked(formula: str) -> None:
    with pytest.raises(ApplicationError):
        safe_formula(formula)


def test_external_reference_is_rewritten_to_local_cache_not_network() -> None:
    assert (
        safe_formula("=VLOOKUP(A1,'[1]Sheet 1'!$A$1:$B$9,2,0)", {(1, "Sheet 1"): "Local cache"})
        == "=VLOOKUP(A1,'Local cache'!$A$1:$B$9,2,0)"
    )


def require_calc() -> None:
    try:
        office_binary()
    except ApplicationError:
        pytest.skip("LibreOffice calculation runtime is not installed")


def test_real_recalculation_updates_cross_sheet_caches_without_rewriting_formulas() -> None:
    require_calc()
    package = WorkbookCells(sample())
    changed = package.patched({package.sheets["Sales"].path: {"B2": ("n", "7", False)}})
    result = recalculate(changed)
    book = load_workbook(io.BytesIO(result), data_only=True)
    assert book["Sales"]["C2"].value == 14
    assert book["Notes"]["D2"].value == 15
    original = load_workbook(io.BytesIO(result), data_only=False)
    assert original["Sales"]["C2"].value == "=B2*2"
    assert original["Notes"]["D2"].value == "=Sales!C2+1"
    assert str(original["Notes"].merged_cells) == "A1:C1"


def test_recalculation_preserves_excel_precision_as_displayed() -> None:
    require_calc()
    book = load_workbook(io.BytesIO(sample()))
    book.calculation.fullPrecision = False
    book["Sales"]["B2"] = 2.14
    book["Sales"]["B2"].number_format = "0.0"
    book["Sales"]["C2"].number_format = "0.0"
    output = io.BytesIO()
    book.save(output)
    result = recalculate(output.getvalue())
    calculated = load_workbook(io.BytesIO(result), data_only=True)
    # Preserve the user's literal input; the workbook's calculation precision is separate.
    assert calculated["Sales"]["B2"].value == 2.14
    assert calculated["Sales"]["C2"].value == 4.2
    assert calculated["Notes"]["D2"].value == 5.2


@pytest.mark.parametrize(("operation", "expected"), [(109, 2), (9, 5)])
def test_recalculation_preserves_hidden_row_subtotal_scope(operation: int, expected: int) -> None:
    require_calc()
    book = load_workbook(io.BytesIO(sample()))
    book["Sales"]["B3"] = 3
    book["Sales"].row_dimensions[3].hidden = True
    book["Sales"].auto_filter.ref = "A1:C3"  # Controls present, but no active filter.
    book["Sales"]["C2"] = f"=SUBTOTAL({operation},B2:B3)"
    output = io.BytesIO()
    book.save(output)
    calculated = load_workbook(io.BytesIO(recalculate(output.getvalue())), data_only=True)
    assert calculated["Sales"]["C2"].value == expected


def test_complex_workbook_api_edits_checks_and_applies_preserving_input(
    db_engine: Engine, tmp_path
) -> None:
    require_calc()
    client = make_client(db_engine, tmp_path)
    try:
        project = str(create_project(db_engine))
        source = client.post(
            f"/api/v1/projects/{project}/sources/excel", files={"file": ("book.xlsx", sample())}
        ).json()["id"]
        update = client.post(
            f"/api/v1/projects/{project}/file-updates",
            json={"source_id": source, "request_id": str(uuid.uuid4())},
        ).json()
        root = f"/api/v1/updates/{update['id']}"
        response = client.get(root + "/editor")
        assert response.status_code == 200, response.text
        page = response.json()
        assert page["workbook"] and not page["can_change_rows"]
        assert len(page["resources"]) == 2
        assert page["rows"][1]["readonly"] == [2]
        assert page["rows"][1]["formulas"] == {"2": "=B2*2"}
        workspace = client.get(root + "/editor/workspace").json()
        assert [s["name"] for s in workspace["sheets"]] == ["Sales", "Notes"]
        assert workspace["sheets"][1]["merges"] == [[1, 1, 3, 1]]
        assert workspace["page"]["limit"] <= 100
        assert workspace["styles"] and workspace["page"]["rows"][0]["style_ids"]
        found = client.get(root + "/editor/find", params={"resource": "Sales", "text": "0001"})
        assert found.json()["cells"] == [{"row": 1, "column": 0}]
        payload = {
            "expected_revision": page["revision"],
            "request_id": str(uuid.uuid4()),
            "resource_id": "Sales",
            "cells": [{"row_id": "r:1", "column": 1, "value": "7"}],
        }
        for patch in [
            {"cells": [{"row_id": "r:1", "column": 2, "value": "9"}]},
            {"append_rows": [str(uuid.uuid4())]},
        ]:
            blocked = client.post(root + "/editor", json={**payload, **patch})
            assert blocked.status_code == 422, blocked.text
        saved = client.post(root + "/editor", json=payload)
        assert saved.status_code == 200, saved.text
        assert client.get(root + "/editor/find", params={"resource": "Sales", "text": "7"}).json()[
            "cells"
        ] == [{"row": 1, "column": 1}]
        assert (
            client.get(
                root + "/editor/find", params={"resource": "missing", "text": "7"}
            ).status_code
            == 404
        )
        assert (
            client.post(root + "/editor", json=payload).json()["revision"]
            == saved.json()["revision"]
        )
        prepared = client.post(
            root + "/editor/prepare",
            json={"expected_revision": saved.json()["revision"], "request_id": str(uuid.uuid4())},
        )
        assert prepared.status_code == 200, prepared.text
        preview = f"/api/v1/sources/{source}/sheets/Sales/preview"
        assert client.get(preview).json()["rows"][0]["col_2"] == 2
        staged = client.get(root + "/editor").json()
        assert staged["rows"][1]["values"] == ["0001", "7", "14"]
        assert (
            client.post(
                root + "/check", json={"expected_revision": prepared.json()["revision"]}
            ).status_code
            == 200
        )
        checked = client.get(f"/api/v1/projects/{project}/updates").json()[0]
        assert checked["check"]["issues"] == []
        applied = client.post(root + "/apply", json={"expected_revision": checked["revision"]})
        assert applied.status_code == 200, applied.text
        assert client.get(preview).json()["rows"][0] == {"col_1": "0001", "col_2": 7, "col_3": 14}
    finally:
        client.close()
        app.dependency_overrides.clear()
