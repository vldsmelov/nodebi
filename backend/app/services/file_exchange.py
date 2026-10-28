"""Source-file round trips reuse the staged update, never chart/export data."""

import csv
import hashlib
import io
import json
import uuid
from pathlib import PurePosixPath
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.errors import ApplicationError
from app.domain.ingestion import ImportProfile
from app.infrastructure.storage import StorageProvider
from app.models import DatasetSnapshot, DataSource, ReportDocument
from app.schemas.file_exchange import FileExchangeRead, FileExchangeStart
from app.services.projects import ProjectService


def source_token(source: DataSource) -> str:
    return hashlib.sha256(
        json.dumps(
            [str(source.id), source.storage_key, source.source_metadata],
            sort_keys=True,
            default=str,
        ).encode()
    ).hexdigest()


def check_exchange(session: Session, row: ReportDocument) -> DataSource | None:
    info = row.payload.get("file_exchange")
    if not info:
        return None
    source = session.scalar(
        select(DataSource)
        .where(
            DataSource.id == uuid.UUID(info["source_id"]), DataSource.project_id == row.project_id
        )
        .execution_options(populate_existing=True)
    )
    if source is None:
        raise ApplicationError("SOURCE_NOT_FOUND", "Data source does not exist.", 404)
    if source_token(source) != info["base_token"]:
        raise ApplicationError(
            "FILE_BASE_CHANGED",
            "The source changed after this file update was started. "
            "Start a new update from the current data.",
            409,
        )
    return source


class FileExchangeService:
    def __init__(self, session: Session, storage: StorageProvider):
        self.session, self.storage = session, storage

    def profiles(self, source: DataSource) -> list[ImportProfile]:
        ids = [uuid.UUID(x) for x in source.source_metadata.get("active_snapshots", {}).values()]
        rows = self.session.scalars(
            select(DatasetSnapshot).where(
                DatasetSnapshot.id.in_(ids), DatasetSnapshot.source_id == source.id
            )
        ).all()
        return [ImportProfile.model_validate(row.profile) for row in rows]

    def template_reason(self, source: DataSource) -> str | None:
        profiles = self.profiles(source)
        if not profiles or source.original_filename.lower().endswith(".xlsm"):
            return (
                "For this workbook, download and edit the current file "
                "to preserve its structure and formulas."
            )
        if any(
            p.options.header_rows != 1
            or p.options.cell_range
            or p.options.skip_data_rows
            or p.options.drop_tail_rows
            for p in profiles
        ):
            return (
                "This source uses a complex range or header. "
                "Download and edit the current file instead."
            )
        return None

    def start(self, project_id: uuid.UUID, request: FileExchangeStart) -> ReportDocument:
        ProjectService(self.session).lock(project_id)
        source = self.session.get(DataSource, request.source_id)
        if source is None or source.project_id != project_id:
            raise ApplicationError("SOURCE_NOT_FOUND", "Data source does not exist.", 404)
        identity = uuid.uuid5(project_id, "file-exchange:" + str(request.request_id))
        existing = self.session.get(ReportDocument, identity)
        if existing is not None:
            if existing.kind != "update" or existing.payload.get("file_exchange", {}).get(
                "source_id"
            ) != str(source.id):
                raise ApplicationError(
                    "FILE_REQUEST_CONFLICT",
                    "This file-update request was already used for another source.",
                    409,
                )
            return existing
        if request.reuse_existing:
            for candidate in self.session.scalars(
                select(ReportDocument)
                .where(
                    ReportDocument.project_id == project_id,
                    ReportDocument.kind == "update",
                    ReportDocument.status.in_(["draft", "checked"]),
                )
                .order_by(ReportDocument.updated_at.desc())
            ):
                info = candidate.payload.get("file_exchange", {})
                if (
                    info.get("source_id") == str(source.id)
                    and info.get("base_token") == source_token(source)
                    and (
                        candidate.payload.get("browser_editor")
                        or candidate.payload.get("browser_commit_receipt")
                        or not candidate.payload.get("items")
                    )
                ):
                    return candidate
        if "connector" in source.source_metadata and not source.source_metadata.get(
            "active_snapshots"
        ):
            raise ApplicationError(
                "IMPORT_INCOMPLETE", "Finish the initial import before replacing its data.", 409
            )
        if not self.storage.exists(source.storage_key):
            raise ApplicationError(
                "FILE_UNAVAILABLE",
                "The source file is unavailable. Restore its storage before updating.",
                409,
            )
        reason = self.template_reason(source)
        info = FileExchangeRead(
            source_id=source.id,
            source_name=source.name,
            connector=source.source_metadata.get("connector", "excel"),
            filename=PurePosixPath(source.original_filename.replace("\\", "/")).name,
            base_token=source_token(source),
            template_available=reason is None,
            template_reason=reason,
        )
        row = ReportDocument(
            id=identity,
            project_id=project_id,
            kind="update",
            payload={
                "name": source.name[:200],
                "items": [],
                "file_exchange": info.model_dump(mode="json"),
            },
        )
        self.session.add(row)
        self.session.commit()
        self.session.refresh(row)
        return row

    def source(self, row: ReportDocument) -> DataSource:
        if row.status not in {"draft", "checked"}:
            raise ApplicationError("UPDATE_CLOSED", "This data update is already closed.", 409)
        source = check_exchange(self.session, row)
        if source is None:
            raise ApplicationError(
                "FILE_UPDATE_REQUIRED", "Start a file update before downloading its files.", 409
            )
        if not self.storage.exists(source.storage_key):
            raise ApplicationError(
                "FILE_UNAVAILABLE",
                "The source file is unavailable. Restore its storage before updating.",
                409,
            )
        return source

    def template(self, source: DataSource) -> tuple[BinaryIO, str]:
        reason = self.template_reason(source)
        if reason:
            raise ApplicationError("TEMPLATE_UNAVAILABLE", reason, 409)
        profiles = self.profiles(source)
        if source.source_metadata.get("connector") == "csv":
            options = profiles[0].options
            text = io.StringIO(newline="")
            with (
                self.storage.open(source.storage_key) as file,
                io.TextIOWrapper(file, encoding=options.encoding, newline="") as incoming,
            ):
                reader = csv.reader(incoming, delimiter=options.delimiter)
                writer = csv.writer(text, delimiter=options.delimiter)
                for _ in range(options.header_row):
                    cells = next(reader, [])
                    if any(cell.lstrip().startswith(("=", "+", "-", "@")) for cell in cells):
                        raise ApplicationError(
                            "TEMPLATE_UNSAFE_HEADER",
                            "CSV headers contain spreadsheet expressions. "
                            "Use the current file through a trusted text importer instead.",
                            409,
                        )
                    writer.writerow(cells)
            return io.BytesIO(text.getvalue().encode(options.encoding)), "template.csv"
        workbook = Workbook()
        workbook.remove(workbook.worksheets[0])
        with self.storage.open(source.storage_key) as file:
            original = load_workbook(file, read_only=True, data_only=True, keep_links=False)
            try:
                for profile in profiles:
                    sheet = workbook.create_sheet(profile.resource_id)
                    header = profile.options.header_row
                    values = next(
                        original[profile.resource_id].iter_rows(
                            min_row=header, max_row=header, values_only=True
                        )
                    )
                    for index, value in enumerate(values, 1):
                        cell = sheet.cell(header, index, str(value) if value is not None else None)
                        if value is not None:
                            cell.data_type = (
                                "s"  # Source labels must not become executable formulas.
                            )
                        field = next(
                            (f for f in profile.fields if f.source_field == f"col_{index}"), None
                        )
                        if field:
                            cell.comment = Comment(
                                f"{field.name}: {field.logical_type.value}; "
                                f"decimal={field.decimal_separator}; date={field.date_format}; "
                                f"scale={field.scale}. Keep this header unchanged.",
                                "Node BI",
                            )
                            # Column defaults avoid creating phantom empty data rows and protect
                            # long identifiers/decimal strings when entering values in Excel.
                            sheet.column_dimensions[get_column_letter(index)].number_format = "@"
                    sheet.freeze_panes = f"A{header + 1}"
            finally:
                original.close()
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return output, "template.xlsx"
