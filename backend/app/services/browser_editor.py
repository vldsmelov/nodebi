"""Paged input editing. Source bytes stay immutable until the normal update is applied."""

import copy
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import InvalidOperation
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.config import Settings
from app.domain.errors import ApplicationError
from app.domain.ingestion import ImportProfile, ReadOptions
from app.domain.workbook_formats import CellFormatPatch
from app.infrastructure.excel import ExcelWorkbookError
from app.infrastructure.storage import StorageProvider
from app.infrastructure.workbook_cells import WorkbookCells, address
from app.infrastructure.workbook_formatting import format_book
from app.infrastructure.worksheet_input_precision import preserve_numbers, table_positions
from app.models import DataSource, ReportDocument
from app.schemas.browser_editor import (
    CellAppearance,
    EditorAction,
    EditorCell,
    EditorColumn,
    EditorMatch,
    EditorPage,
    EditorRow,
    EditorSave,
    EditorSearch,
    EditorWorkspace,
    SheetAppearance,
)
from app.schemas.reporting import UpdateItem, UpdateRead, UpdateSave
from app.services.browser_inputs import legacy_profiles
from app.services.file_exchange import FileExchangeService
from app.services.normalization import convert_input_value
from app.services.projects import ProjectService
from app.services.source_updates import SourceUpdateService, change_document, read_update


def unsupported() -> ApplicationError:
    return ApplicationError(
        "BROWSER_EDITOR_UNAVAILABLE",
        "Browser editing supports plain CSV and XLSX tables only. "
        "For formulas, complex workbooks or large Excel files, "
        "download and edit the original file.",
        409,
    )


class BrowserEditorService:
    def __init__(self, session: Session, storage: StorageProvider):
        self.session, self.storage = session, storage
        self.exchange = FileExchangeService(session, storage)

    def package(self, key: str) -> WorkbookCells:
        if not hasattr(self, "_packages"):
            self._packages: dict[str, WorkbookCells] = {}
        if key not in self._packages:
            with self.storage.open(key) as file:
                self._packages[key] = WorkbookCells(file.read())
        return self._packages[key]

    def context(self, row: ReportDocument) -> tuple[DataSource, str, list[ImportProfile]]:
        source = self.exchange.source(row)
        legacy = "connector" not in source.source_metadata
        if legacy:
            workbook = self.open_workbook(source.storage_key)
            workbook.close()
        if not legacy and self.exchange.template_reason(source):
            raise unsupported()
        items = read_update(row).items
        try:
            profiles = (
                legacy_profiles(self.session, self.storage, source)
                if legacy
                else self.exchange.profiles(source)
            )
        except (ValueError, ExcelWorkbookError) as error:
            raise unsupported() from error
        key = source.storage_key
        if items:
            version = SourceUpdateService(self.session, self.storage).version(
                row.project_id, items[0]
            )
            key = version.storage_key
            if source.source_metadata.get("connector") == "csv" and (
                ReadOptions.model_validate(version.workbook.get("read_options", {}))
                != profiles[0].options
            ):
                raise unsupported()
            for binding in items[0].dataset_bindings:
                profile = next((p for p in profiles if p.resource_id == binding.resource_id), None)
                if (
                    profile is None
                    or binding.resource_id != binding.target_resource_id
                    or binding.options is not None
                    and binding.options != profile.options
                ):
                    raise unsupported()
        if key.lower().endswith(".xlsm"):
            raise unsupported()
        state = row.payload.get("browser_editor", {})
        if state and state["storage_key"] != key:
            raise ApplicationError(
                "EDITOR_INPUT_CHANGED",
                "The selected replacement changed. "
                "Discard the browser draft before editing another file.",
                409,
            )
        return source, key, profiles

    def raw(
        self, source: DataSource, key: str, profile: ImportProfile
    ) -> tuple[list[str], list[list[str]]]:
        options = profile.options
        if source.source_metadata.get("connector") == "csv":
            with (
                self.storage.open(key) as file,
                io.TextIOWrapper(file, encoding=options.encoding, newline="") as text,
            ):
                reader = csv.reader(text, delimiter=options.delimiter, strict=True)
                for _ in range(options.header_row - 1):
                    next(reader, None)
                headers = next(reader, [])
                if not headers:
                    raise unsupported()
                records = []
                for values in reader:
                    if len(values) != len(headers):
                        raise ApplicationError(
                            "INVALID_IMPORT", "The file has inconsistent column counts."
                        )
                    records.append(values)
                    if len(records) > 250000 or len(records) * len(headers) > 2000000:
                        raise unsupported()
            return headers, records
        workbook = self.open_workbook(key)
        try:
            if profile.resource_id not in workbook.sheetnames:
                raise unsupported()
            sheet = workbook[profile.resource_id]
            headers = [self.text(c.value) for c in sheet[options.header_row]]
            records = (
                [
                    [self.text(c.value) for c in cells]
                    for cells in sheet.iter_rows(
                        min_row=options.header_row + 1,
                        max_row=sheet.max_row,
                        max_col=len(headers),
                    )
                ]
                if sheet.max_row > options.header_row
                else []
            )
            return headers, records
        finally:
            workbook.close()

    def open_workbook(self, key: str) -> Any:
        if key.lower().endswith(".xlsm"):
            raise unsupported()
        with self.storage.open(key) as file:
            probe = load_workbook(file, read_only=True, data_only=False, keep_links=False)
            try:
                if sum((s.max_row or 1) * (s.max_column or 1) for s in probe.worksheets) > 200000:
                    raise unsupported()
            finally:
                probe.close()
            file.seek(0)
            workbook = load_workbook(file, read_only=False, data_only=False, keep_links=True)
            try:
                self.validate_workbook(workbook)
            except Exception:
                workbook.close()
                raise
            return workbook

    @staticmethod
    def text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def validate_workbook(workbook: Any) -> None:
        # No formula evaluator is introduced. Saving a formula workbook with openpyxl
        # invalidates cached results even when only another sheet is edited.
        if workbook.defined_names or workbook._external_links:
            raise unsupported()
        total = 0
        for sheet in workbook.worksheets:
            total += sheet.max_row * sheet.max_column
            if (
                total > 200000
                or sheet.tables
                or sheet.merged_cells
                or sheet._charts
                or sheet._images
                or sheet.data_validations.count
                or sheet.conditional_formatting
            ):
                raise unsupported()
            if any(cell.data_type == "f" for cells in sheet.iter_rows() for cell in cells):
                raise unsupported()

    def page(
        self, row: ReportDocument, resource: str | None, offset: int, limit: int
    ) -> EditorPage:
        source, key, profiles = self.context(row)
        profile = next(
            (p for p in profiles if p.resource_id == (resource or profiles[0].resource_id)), None
        )
        if profile is None:
            raise ApplicationError(
                "EDITOR_RESOURCE_NOT_FOUND", "Select an available input table.", 404
            )
        headers, records = self.raw(source, key, profile)
        state = (
            row.payload.get("browser_editor", {}).get("resources", {}).get(profile.resource_id, {})
        )
        deleted = set(state.get("deleted", []))
        ids = [f"r:{i}" for i in range(len(records))] + state.get("added", [])
        ids = [identity for identity in ids if identity not in deleted]
        limit = max(1, min(limit, 100, 2000 // max(1, len(headers))))
        offset = min(offset, max(0, ((len(ids) - 1) // limit) * limit))
        rows = []
        for identity in ids[offset : offset + limit]:
            values = (
                list(records[int(identity[2:])])
                if identity.startswith("r:")
                else [""] * len(headers)
            )
            for column, value in state.get("cells", {}).get(identity, {}).items():
                values[int(column)] = value
            record = EditorRow(id=identity, values=values)
            record.formats = {
                int(c): CellFormatPatch.model_validate(patch)
                for c, patch in state.get("formats", {}).get(identity, {}).items()
            }
            if source.source_metadata.get("connector") != "csv":
                sheet = self.package(key).sheets[profile.resource_id]
                physical_row = (
                    profile.options.header_row + int(identity[2:])
                    if identity.startswith("r:")
                    else -1
                )
                record.style_ids = {
                    c: sheet.style_ids.get(address(physical_row, c), 0) for c in range(len(headers))
                }
            rows.append(record)
        fields = {f.source_field: f for f in profile.fields}
        columns = []
        for index, header in enumerate(headers):
            field = fields.get(f"col_{index + 1}")
            columns.append(
                EditorColumn(
                    index=index,
                    name=header or f"Column {index + 1}",
                    description=field.logical_type.value if field else "Not used in import",
                    date_format=field.date_format
                    if field and field.logical_type.value in {"date", "datetime"}
                    else None,
                    decimal_separator=field.decimal_separator
                    if field and field.logical_type.value in {"decimal", "float"}
                    else None,
                )
            )
        return EditorPage(
            revision=row.revision,
            resource_id=profile.resource_id,
            resources=[p.resource_id for p in profiles],
            columns=columns,
            rows=rows,
            can_format=source.source_metadata.get("connector") != "csv",
            total_rows=len(ids),
            offset=offset,
            limit=limit,
            pending=bool(row.payload.get("browser_editor")),
        )

    def workspace(self, row: ReportDocument) -> EditorWorkspace:
        first = self.page(row, None, 0, 100)
        pages = [
            first if name == first.resource_id else self.page(row, name, 0, 1)
            for name in first.resources
        ]
        return EditorWorkspace(
            page=first,
            styles=[
                CellAppearance.model_validate(s)
                for s in self.package(self.context(row)[1]).display_styles
            ]
            if first.can_format
            else [],
            sheets=[
                SheetAppearance(
                    name=page.resource_id, rows=page.total_rows, columns=len(page.columns)
                )
                for page in pages
            ],
        )

    def find(self, row: ReportDocument, resource: str, text: str) -> EditorSearch:
        source, key, profiles = self.context(row)
        profile = next((p for p in profiles if p.resource_id == resource), None)
        if profile is None:
            raise ApplicationError(
                "EDITOR_RESOURCE_NOT_FOUND", "Select an available input table.", 404
            )
        _, records = self.raw(source, key, profile)
        state = row.payload.get("browser_editor", {}).get("resources", {}).get(resource, {})
        ids = [f"r:{i}" for i in range(len(records))] + state.get("added", [])
        deleted = set(state.get("deleted", []))
        matches = []
        index = 0
        for identity_row in ids:
            if identity_row in deleted:
                continue
            values = (
                records[int(identity_row[2:])]
                if identity_row.startswith("r:")
                else [""] * len(records[0] if records else profile.fields)
            )
            for column, original in enumerate(values):
                value = state.get("cells", {}).get(identity_row, {}).get(str(column), original)
                if text.casefold() in value.casefold():
                    matches.append(EditorMatch(row=index, column=column))
                    if len(matches) >= 100:
                        return EditorSearch(cells=matches)
            index += 1
        return EditorSearch(cells=matches)

    def data_shape(self, source: DataSource, key: str, profile: ImportProfile) -> tuple[int, int]:
        headers, records = self.raw(source, key, profile)
        return len(records), len(headers)

    def save(self, row: ReportDocument, request: EditorSave, offset: int, limit: int) -> EditorPage:
        ProjectService(self.session).lock(row.project_id)
        self.session.refresh(row)
        # Keep existing value-only retry receipts compatible across this additive contract change.
        digest = hashlib.sha256(
            request.model_dump_json(exclude={"formats"} if not request.formats else None).encode()
        ).hexdigest()
        receipt = row.payload.get("browser_save_receipt", {})
        if receipt.get("request_id") == str(request.request_id):
            if receipt.get("digest") != digest or row.revision != receipt["revision"]:
                raise ApplicationError(
                    "REPORT_CONFLICT", "The report changed. Reload before applying your draft.", 409
                )
            return self.page(row, request.resource_id, offset, limit)
        if row.revision != request.expected_revision:
            raise ApplicationError(
                "REPORT_CONFLICT", "The report changed. Reload before applying your draft.", 409
            )
        source, key, profiles = self.context(row)
        profile = next((p for p in profiles if p.resource_id == request.resource_id), None)
        if profile is None:
            raise unsupported()
        row_count, column_count = self.data_shape(source, key, profile)
        self.validate_edits(key, profile, request)
        if len(request.cells) + len(request.formats) > 2000:
            raise ApplicationError(
                "EDITOR_CHANGE_LIMIT", "The pasted range is too large. Paste a smaller range."
            )
        self.validate_formats(key, profile, request)
        editor = copy.deepcopy(
            row.payload.get("browser_editor") or {"storage_key": key, "resources": {}}
        )
        resource = editor["resources"].setdefault(
            profile.resource_id, {"cells": {}, "deleted": [], "added": []}
        )
        resource["added"] = list(
            dict.fromkeys([*resource["added"], *("a:" + str(x) for x in request.append_rows)])
        )
        valid = {f"r:{i}" for i in range(row_count)} | set(resource["added"])
        if any(
            identity not in valid
            for identity in [
                *request.delete_rows,
                *request.restore_rows,
                *(c.row_id for c in request.cells),
                *(c.row_id for c in request.formats),
            ]
        ):
            raise ApplicationError(
                "EDITOR_ROW_NOT_FOUND", "The edited row is no longer available.", 409
            )
        for cell in request.cells:
            if cell.column >= column_count:
                raise ApplicationError(
                    "EDITOR_COLUMN_NOT_FOUND", "The edited column is no longer available.", 409
                )
            if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", cell.value):
                raise ApplicationError(
                    "EDITOR_INVALID_TEXT", "Remove unsupported control characters from the cell."
                )
            if source.source_metadata.get("connector") == "csv":
                try:
                    cell.value.encode(profile.options.encoding)
                except UnicodeEncodeError as error:
                    raise ApplicationError(
                        "EDITOR_ENCODING",
                        "The cell contains characters unsupported by this CSV encoding. "
                        "Use compatible characters or import a UTF-8 file.",
                    ) from error
            resource["cells"].setdefault(cell.row_id, {})[str(cell.column)] = cell.value
            self.record_cell_options(resource, cell)
        for change in request.formats:
            if change.column >= column_count:
                raise ApplicationError(
                    "EDITOR_COLUMN_NOT_FOUND", "The edited column is no longer available.", 409
                )
            resource.setdefault("formats", {}).setdefault(change.row_id, {})[str(change.column)] = (
                change.patch.model_dump(exclude_none=True)
            )
            resource.setdefault("format_options", {}).setdefault(change.row_id, {})[
                str(change.column)
            ] = {"override_protection": change.override_protection}
        resource["deleted"] = sorted(
            (set(resource["deleted"]) | set(request.delete_rows)) - set(request.restore_rows)
        )
        total_cells = sum(
            sum(len(v) for v in r["cells"].values())
            + sum(len(v) for v in r.get("formats", {}).values())
            for r in editor["resources"].values()
        )
        total_ops = sum(len(r["added"]) + len(r["deleted"]) for r in editor["resources"].values())
        if (
            total_cells > 10000
            or total_ops > 10000
            or row_count + len(resource["added"]) > 250000
            or (row_count + len(resource["added"])) * column_count > 2000000
            or len(json.dumps(editor).encode()) > 4000000
        ):
            raise ApplicationError(
                "EDITOR_CHANGE_LIMIT",
                "The browser draft is too large. "
                "Prepare this update file before making more changes.",
                409,
            )
        payload = {k: v for k, v in row.payload.items() if k not in {"check", "prepared"}}
        payload.update(
            browser_editor=editor,
            browser_save_receipt={
                "request_id": str(request.request_id),
                "digest": digest,
                "revision": row.revision + 1,
            },
        )
        change_document(self.session, row, request.expected_revision, payload, "draft")
        self.session.commit()
        return self.page(row, request.resource_id, offset, limit)

    def validate_edits(self, key: str, profile: ImportProfile, request: EditorSave) -> None:
        if any(c.formula or c.override_protection for c in request.cells):
            raise ApplicationError(
                "EDITOR_WORKBOOK_MODE_REQUIRED",
                "Formula editing is available in worksheet mode, not in an imported table.",
            )

    def record_cell_options(self, resource: dict[str, Any], cell: EditorCell) -> None:
        """Source-specific metadata for the explicit, typed cell operation."""

    def validate_formats(self, key: str, profile: ImportProfile, request: EditorSave) -> None:
        if request.formats and not key.lower().endswith(".xlsx"):
            raise ApplicationError(
                "EDITOR_FORMAT_UNAVAILABLE", "CSV files do not store cell formatting."
            )

    def discard(self, row: ReportDocument, revision: int) -> UpdateRead:
        ProjectService(self.session).lock(row.project_id)
        self.session.refresh(row)
        if row.status not in {"draft", "checked"}:
            raise ApplicationError("UPDATE_CLOSED", "This data update is already closed.", 409)
        payload = {
            k: v
            for k, v in row.payload.items()
            if k not in {"browser_editor", "browser_save_receipt"}
        }
        change_document(self.session, row, revision, payload, "draft")
        self.session.commit()
        return read_update(row)

    def prepare_metadata(self, output: io.BytesIO) -> dict[str, Any]:
        return {}

    def commit(self, row: ReportDocument, request: EditorAction, settings: Settings) -> UpdateRead:
        ProjectService(self.session).lock(row.project_id)
        self.session.refresh(row)
        receipt = row.payload.get("browser_commit_receipt", {})
        if receipt.get("request_id") == str(request.request_id):
            if (
                receipt.get("expected_revision") != request.expected_revision
                or row.revision != receipt["revision"]
            ):
                raise ApplicationError(
                    "REPORT_CONFLICT", "The report changed. Reload before applying your draft.", 409
                )
            return read_update(row)
        if row.revision != request.expected_revision:
            raise ApplicationError(
                "REPORT_CONFLICT", "The report changed. Reload before applying your draft.", 409
            )
        source, key, profiles = self.context(row)
        editor = row.payload.get("browser_editor")
        if not editor:
            raise ApplicationError("EDITOR_EMPTY", "Edit and save a browser draft first.", 409)
        output, filename = self.render(source, key, profiles, editor)
        return self.stage(row, request, settings, source, profiles, output, filename)

    def render(
        self, source: DataSource, key: str, profiles: list[ImportProfile], editor: dict[str, Any]
    ) -> tuple[io.BytesIO, str]:
        if source.source_metadata.get("connector") != "csv" and not any(
            state.get("cells") or state.get("added") or state.get("deleted")
            for state in editor["resources"].values()
        ):
            # Styling alone must not reserialize even one existing value through openpyxl.
            source_formats = {
                p.resource_id: {
                    address(
                        p.options.header_row + int(identity[2:]), int(column)
                    ): CellFormatPatch.model_validate(patch)
                    for identity, cells in editor["resources"]
                    .get(p.resource_id, {})
                    .get("formats", {})
                    .items()
                    for column, patch in cells.items()
                }
                for p in profiles
            }
            return io.BytesIO(
                format_book(self.package(key).content, source_formats)
            ), "browser-edited.xlsx"
        output = io.BytesIO()
        if source.source_metadata.get("connector") == "csv":
            profile = profiles[0]
            headers, records = self.raw(source, key, profile)
            state = editor["resources"].get(profile.resource_id, {})
            text = io.StringIO(newline="")
            writer = csv.writer(text, delimiter=profile.options.delimiter)
            with (
                self.storage.open(key) as file,
                io.TextIOWrapper(file, encoding=profile.options.encoding, newline="") as incoming,
            ):
                reader = csv.reader(incoming, delimiter=profile.options.delimiter)
                for _ in range(profile.options.header_row):
                    writer.writerow(next(reader))
            ids = [f"r:{i}" for i in range(len(records))] + state.get("added", [])
            deleted = set(state.get("deleted", []))
            for identity in ids:
                if identity in deleted:
                    continue
                values = (
                    list(records[int(identity[2:])])
                    if identity.startswith("r:")
                    else [""] * len(headers)
                )
                for column, value in state.get("cells", {}).get(identity, {}).items():
                    values[int(column)] = value
                writer.writerow(values)
            output.write(text.getvalue().encode(profile.options.encoding))
            filename = "browser-edited.csv"
        else:
            counts: dict[str, int] = {}
            with self.storage.open(key) as file:
                workbook = load_workbook(file, data_only=False, keep_links=True)
                try:
                    self.validate_workbook(workbook)
                    counts = {
                        p.resource_id: max(
                            0, workbook[p.resource_id].max_row - p.options.header_row
                        )
                        for p in profiles
                    }
                    for profile in profiles:
                        state = editor["resources"].get(profile.resource_id)
                        if not state:
                            continue
                        sheet = workbook[profile.resource_id]
                        start = profile.options.header_row + 1
                        for identity, cells in state["cells"].items():
                            if identity.startswith("r:"):
                                for column, value in cells.items():
                                    cell = sheet.cell(start + int(identity[2:]), int(column) + 1)
                                    cell.value = self.cell_value(
                                        source, profile, int(column), value
                                    )
                                    if isinstance(cell.value, str):
                                        cell.data_type = "s"
                        for identity in sorted(
                            (x for x in state["deleted"] if x.startswith("r:")),
                            key=lambda x: int(x[2:]),
                            reverse=True,
                        ):
                            sheet.delete_rows(start + int(identity[2:]))
                        for identity in state["added"]:
                            if identity in state["deleted"]:
                                continue
                            index = sheet.max_row + 1
                            for column in range(sheet.max_column):
                                value = state["cells"].get(identity, {}).get(str(column), "")
                                cell = sheet.cell(
                                    index,
                                    column + 1,
                                    self.cell_value(source, profile, column, value),
                                )
                                if isinstance(cell.value, str):
                                    cell.data_type = "s"
                    self.validate_workbook(workbook)
                    workbook.save(output)
                finally:
                    workbook.close()
            filename = "browser-edited.xlsx"
            output = io.BytesIO(
                preserve_numbers(
                    self.package(key),
                    output.getvalue(),
                    profiles,
                    editor["resources"],
                    counts,
                    lambda p, c, v: self.cell_value(source, p, c, v),
                )
            )
            formats: dict[str, dict[str, CellFormatPatch]] = {}
            for profile in profiles:
                state = editor["resources"].get(profile.resource_id, {})
                if not state.get("formats"):
                    continue
                for identity, position in table_positions(
                    counts[profile.resource_id], profile, state
                ).items():
                    for column, patch in state.get("formats", {}).get(identity, {}).items():
                        formats.setdefault(profile.resource_id, {})[
                            address(position, int(column))
                        ] = CellFormatPatch.model_validate(patch)
            output = io.BytesIO(format_book(output.getvalue(), formats))
        return output, filename

    def stage(
        self,
        row: ReportDocument,
        request: EditorAction,
        settings: Settings,
        source: DataSource,
        profiles: list[ImportProfile],
        output: io.BytesIO,
        filename: str,
    ) -> UpdateRead:
        if output.getbuffer().nbytes > settings.max_upload_bytes:
            raise unsupported()
        output.seek(0)
        service = SourceUpdateService(self.session, self.storage)
        # Upload stages an immutable version, not active data. If the following CAS fails,
        # the version stays available without overwriting the source or browser draft.
        source_id = source.id
        # The uploader reserves artifacts in a separate transaction. Do not hold a
        # project row lock across that FK reservation (PostgreSQL would deadlock).
        self.session.commit()
        version = service.upload_version(source_id, filename, output, settings, profiles[0].options)
        ProjectService(self.session).lock(row.project_id)
        self.session.refresh(row)
        if row.revision != request.expected_revision:
            raise ApplicationError(
                "REPORT_CONFLICT", "The report changed. Reload before applying your draft.", 409
            )
        row.payload = {
            k: v
            for k, v in row.payload.items()
            if k not in {"browser_editor", "browser_save_receipt"}
        }
        row.payload = {
            **row.payload,
            "browser_commit_receipt": {
                "request_id": str(request.request_id),
                "expected_revision": request.expected_revision,
                "revision": row.revision + 1,
                "version_id": str(version.id),
                **self.prepare_metadata(output),
            },
        }
        service.save(
            row,
            UpdateSave(
                expected_revision=row.revision,
                name=row.payload.get("name", source.name),
                items=[
                    read_update(row).items[0].model_copy(update={"version_id": version.id})
                    if read_update(row).items
                    else UpdateItem(source_id=source.id, version_id=version.id)
                ],
            ),
        )
        return read_update(row)

    @staticmethod
    def cell_value(source: DataSource, profile: ImportProfile, column: int, value: str) -> Any:
        if "connector" in source.source_metadata:
            return value or None
        field = next((f for f in profile.fields if f.source_field == f"col_{column + 1}"), None)
        if field is None:
            return value or None
        try:
            return convert_input_value(value if value else None, field)
        except (ValueError, InvalidOperation, OverflowError) as error:
            raise ApplicationError(
                "EDITOR_VALUE_TYPE",
                f"Invalid value for {field.name}: {str(error) or 'check the field type.'}",
            ) from error


def browser_editor_service(
    session: Session, storage: StorageProvider, row: ReportDocument
) -> BrowserEditorService:
    service = BrowserEditorService(session, storage)
    try:
        source, key, _ = service.context(row)
        if source.source_metadata.get("connector") != "csv":
            book = service.open_workbook(key)
            book.close()
    except ApplicationError as error:
        source = service.exchange.source(row)
        if error.code != "BROWSER_EDITOR_UNAVAILABLE" or not source.storage_key.lower().endswith(
            ".xlsx"
        ):
            raise
        from app.services.workbook_editor import WorkbookEditorService

        return WorkbookEditorService(session, storage)
    return service
