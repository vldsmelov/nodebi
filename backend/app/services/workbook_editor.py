"""Edit physical input cells of complex workbooks using existing staged-update semantics."""

import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.utils.datetime import MAC_EPOCH, to_excel

from app.domain.datasets import LogicalType
from app.domain.errors import ApplicationError
from app.domain.ingestion import ImportField, ImportProfile
from app.domain.workbook_formats import CellFormatPatch
from app.infrastructure.spreadsheet_recalculation import recalculate
from app.infrastructure.workbook_cells import NS, WorkbookCells, address
from app.infrastructure.workbook_formatting import format_book
from app.infrastructure.workbook_formula_editing import authored_formula, invalidate_formula_chain
from app.models import DataSource, ReportDocument
from app.schemas.browser_editor import (
    CellAppearance,
    EditorCell,
    EditorColumn,
    EditorPage,
    EditorRow,
    EditorSave,
    EditorWorkspace,
    SheetAppearance,
)
from app.services.browser_editor import BrowserEditorService, unsupported
from app.services.source_updates import SourceUpdateService, read_update


class WorkbookEditorService(BrowserEditorService):
    def context(self, row: ReportDocument) -> tuple[DataSource, str, list[ImportProfile]]:
        source = self.exchange.source(row)
        if source.source_metadata.get("periods"):
            raise unsupported()
        items = read_update(row).items
        key = (
            SourceUpdateService(self.session, self.storage)
            .version(row.project_id, items[0])
            .storage_key
            if items
            else source.storage_key
        )
        if not key.lower().endswith(".xlsx"):
            raise unsupported()
        state = row.payload.get("browser_editor", {})
        if state and state["storage_key"] != key:
            raise ApplicationError(
                "EDITOR_INPUT_CHANGED",
                "The selected replacement changed. "
                "Discard the browser draft before editing another file.",
                409,
            )
        package = self.package(key)
        profiles = [
            ImportProfile(
                resource_id=name,
                name=name[:200],
                fields=[
                    ImportField(
                        field_id="cell",
                        source_field="col_1",
                        name="Cell",
                        logical_type=LogicalType.STRING,
                    )
                ],
            )
            for name in package.sheets
        ]
        return source, key, profiles

    def data_shape(self, source: DataSource, key: str, profile: ImportProfile) -> tuple[int, int]:
        sheet = self.package(key).sheets[profile.resource_id]
        return sheet.rows, sheet.columns

    def raw(
        self, source: DataSource, key: str, profile: ImportProfile
    ) -> tuple[list[str], list[list[str]]]:
        package = self.package(key)
        sheet = package.sheets[profile.resource_id]
        return (
            [get_column_letter(c + 1) for c in range(sheet.columns)],
            [
                [package.value(sheet.cells.get(address(r, c))) for c in range(sheet.columns)]
                for r in range(sheet.rows)
            ],
        )

    def page(
        self, row: ReportDocument, resource: str | None, offset: int, limit: int
    ) -> EditorPage:
        if resource is None:
            _, key, _ = self.context(row)
            resource = next(
                (
                    sheet.name
                    for sheet in self.package(key).sheets.values()
                    if not sheet.protected
                    and any(not sheet.readonly(coordinate) for coordinate in sheet.cells)
                ),
                None,
            )
        _, key, profiles = self.context(row)
        package = self.package(key)
        resource = resource or profiles[0].resource_id
        if resource not in package.sheets:
            raise ApplicationError(
                "EDITOR_RESOURCE_NOT_FOUND", "Select an available input table.", 404
            )
        sheet = package.sheets[resource]
        limit = min(limit, 100, max(1, 2000 // sheet.columns))
        offset = min(offset, max(0, ((sheet.rows - 1) // limit) * limit))
        state = row.payload.get("browser_editor", {}).get("resources", {}).get(resource, {})
        records = []
        for index in range(offset, min(sheet.rows, offset + limit)):
            identity = f"r:{index}"
            values = [
                package.value(sheet.cells.get(address(index, c))) for c in range(sheet.columns)
            ]
            for column, value in state.get("cells", {}).get(identity, {}).items():
                values[int(column)] = value
            record = EditorRow(id=identity, values=values)
            index = int(record.id[2:])
            record.readonly = [c for c in range(sheet.columns) if sheet.readonly(address(index, c))]
            record.style_ids = {
                c: sheet.style_ids.get(address(index, c), 0) for c in range(sheet.columns)
            }
            record.formats = {
                int(c): CellFormatPatch.model_validate(patch)
                for c, patch in state.get("formats", {}).get(identity, {}).items()
            }
            record.format_readonly = [
                c
                for c in range(sheet.columns)
                if sheet.structural_readonly(address(index, c))
                or (sheet.protected and address(index, c) not in sheet.unlocked)
            ]
            record.number_values = {
                c: sheet.cells[at].findtext(NS + "v", "")
                for c in range(sheet.columns)
                if (at := address(index, c)) in sheet.cells
                and sheet.cells[at].get("t", "n") == "n"
                and sheet.cells[at].findtext(NS + "v", "") != ""
            }
            record.formulas = {
                c: sheet.shared_formulas[address(index, c)][1]
                if address(index, c) in sheet.shared_formulas
                else "=" + (sheet.cells[address(index, c)].findtext(NS + "f") or "")
                for c in record.readonly
                if address(index, c) in sheet.cells
                and sheet.cells[address(index, c)].find(NS + "f") is not None
            }
            record.editable_in_workbook_mode = [
                c for c in record.readonly if not sheet.structural_readonly(address(index, c))
            ]
            for column, value in state.get("cells", {}).get(identity, {}).items():
                c = int(column)
                record.number_values.pop(c, None)
                if state.get("cell_options", {}).get(identity, {}).get(column, {}).get("formula"):
                    record.formulas[c] = value
                    record.pending_formulas.append(c)
                    record.values[c] = ""  # No stale numeric result for an uncalculated formula.
                else:
                    record.formulas.pop(c, None)
            records.append(record)
        return EditorPage(
            revision=row.revision,
            resource_id=resource,
            resources=list(package.sheets),
            columns=[
                EditorColumn(index=c, name=get_column_letter(c + 1), description="Worksheet column")
                for c in range(sheet.columns)
            ],
            rows=records,
            total_rows=sheet.rows,
            offset=offset,
            limit=limit,
            pending=bool(row.payload.get("browser_editor")),
            workbook=True,
            can_change_rows=False,
            can_format=True,
            external_cached_values=bool(package.external),
            precision_as_displayed=package.precision_as_displayed,
            date_1904=package.epoch == MAC_EPOCH,
        )

    def workspace(self, row: ReportDocument) -> EditorWorkspace:
        _, key, _ = self.context(row)
        package = self.package(key)
        return EditorWorkspace(
            page=self.page(row, None, 0, 100),
            styles=[CellAppearance.model_validate(s) for s in package.display_styles],
            sheets=[
                SheetAppearance(
                    name=s.name,
                    rows=s.rows,
                    columns=s.columns,
                    widths=s.widths,
                    heights=s.heights,
                    merges=s.merges,
                )
                for s in package.sheets.values()
            ],
        )

    def validate_edits(self, key: str, profile: ImportProfile, request: EditorSave) -> None:
        if request.append_rows or request.delete_rows or request.restore_rows:
            raise ApplicationError(
                "EDITOR_STRUCTURE_PROTECTED",
                "Workbook rows are protected to preserve formulas and references. "
                "Edit existing input cells.",
            )
        sheet = self.package(key).sheets[profile.resource_id]
        for cell in request.cells:
            if not re.fullmatch(r"r:[0-9]+", cell.row_id):
                raise ApplicationError(
                    "EDITOR_ROW_NOT_FOUND", "The edited row is no longer available.", 409
                )
            coordinate = address(int(cell.row_id[2:]), cell.column)
            if sheet.structural_readonly(coordinate):
                raise ApplicationError(
                    "EDITOR_STRUCTURE_PROTECTED",
                    "Edit the top-left cell of a merged area. "
                    "Array and data-table formulas stay protected.",
                )
            if sheet.readonly(coordinate) and not cell.override_protection:
                raise ApplicationError(
                    "EDITOR_CELL_READONLY",
                    "Enable Edit calculations to change formulas or protected cell values.",
                )
            if cell.formula:
                cell.value = authored_formula(cell.value, set(self.package(key).external))

    def record_cell_options(self, resource: dict[str, Any], cell: EditorCell) -> None:
        resource.setdefault("cell_options", {}).setdefault(cell.row_id, {})[str(cell.column)] = {
            "formula": cell.formula,
            "override_protection": cell.override_protection,
        }

    def validate_formats(self, key: str, profile: ImportProfile, request: EditorSave) -> None:
        sheet = self.package(key).sheets[profile.resource_id]
        for change in request.formats:
            if not re.fullmatch(r"r:[0-9]+", change.row_id):
                raise ApplicationError(
                    "EDITOR_ROW_NOT_FOUND", "The edited row is no longer available.", 409
                )
            coordinate = address(int(change.row_id[2:]), change.column)
            if sheet.structural_readonly(coordinate):
                raise ApplicationError(
                    "EDITOR_STRUCTURE_PROTECTED",
                    "Edit the top-left cell of a merged area. "
                    "Array and data-table formulas stay protected.",
                )
            if (
                sheet.protected
                and coordinate not in sheet.unlocked
                and not change.override_protection
            ):
                raise ApplicationError(
                    "EDITOR_CELL_READONLY",
                    "Enable Edit calculations to change formulas or protected cell values.",
                )

    def render(
        self, source: DataSource, key: str, profiles: list[ImportProfile], editor: dict[str, Any]
    ) -> tuple[io.BytesIO, str]:
        package = self.package(key)
        updates: dict[str, dict[str, tuple[str, str, bool]]] = {}
        changed_formulas = False
        for name, state in editor["resources"].items():
            sheet = package.sheets[name]
            if state.get("added") or state.get("deleted"):
                raise ApplicationError(
                    "EDITOR_STRUCTURE_PROTECTED",
                    "Workbook rows are protected to preserve formulas and references. "
                    "Edit existing input cells.",
                )
            for identity, cells in state["cells"].items():
                for column, value in cells.items():
                    coordinate = address(int(identity[2:]), int(column))
                    options = state.get("cell_options", {}).get(identity, {}).get(column, {})
                    if sheet.structural_readonly(coordinate) or (
                        sheet.readonly(coordinate) and not options.get("override_protection")
                    ):
                        raise ApplicationError(
                            "EDITOR_CELL_READONLY",
                            "Enable Edit calculations to change formulas or protected cell values.",
                        )
                    if options.get("formula"):
                        changed_formulas = True
                        value = authored_formula(value, set(package.external))
                        updates.setdefault(sheet.path, {})[coordinate] = ("formula", value, False)
                        continue
                    original = sheet.cells.get(coordinate)
                    was_formula = original is not None and original.find(NS + "f") is not None
                    changed_formulas |= was_formula
                    kind = (
                        original.get("t", "n")
                        if original is not None and not was_formula
                        else "inlineStr"
                    )
                    output_kind = "inlineStr"
                    if value and kind == "b":
                        if value.lower() not in {"true", "false", "1", "0"}:
                            raise ApplicationError(
                                "EDITOR_VALUE_TYPE", f"Enter true or false at {name}!{coordinate}."
                            )
                        output_kind, value = "b", "1" if value.lower() in {"true", "1"} else "0"
                    elif value and kind == "n":
                        try:
                            style = int(original.get("s", "0")) if original is not None else 0
                            if style < len(package.formats) and is_date_format(
                                package.formats[style]
                            ):
                                value = str(
                                    to_excel(  # type: ignore[no-untyped-call]
                                        datetime.fromisoformat(value), package.epoch
                                    )
                                )
                            else:
                                number = Decimal(
                                    value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
                                )
                                if not number.is_finite():
                                    raise ValueError("Not finite")
                                value = str(number)
                            output_kind = "n"
                        except (ValueError, InvalidOperation) as error:
                            raise ApplicationError(
                                "EDITOR_VALUE_TYPE",
                                f"Enter a valid number or ISO date at {name}!{coordinate}.",
                            ) from error
                    elif (original is None or was_formula) and re.fullmatch(
                        r"-?(?:0|[1-9]\d*)(?:[.,]\d+)?", value
                    ):
                        output_kind, value = "n", value.replace(",", ".")
                    updates.setdefault(sheet.path, {})[coordinate] = (output_kind, value, False)
        # Shared formulas are a storage optimization, not an inseparable array operation.
        # Expand only affected groups, preserving the translated expression of every other member.
        for sheet in package.sheets.values():
            changes = updates.get(sheet.path, {})
            groups = {
                group
                for coordinate, (group, _) in sheet.shared_formulas.items()
                if coordinate in changes
            }
            for coordinate, (group, formula) in sheet.shared_formulas.items():
                if group in groups and coordinate not in changes:
                    changes[coordinate] = ("formula", formula, False)
        formats = {
            name: {
                address(int(identity[2:]), int(column)): CellFormatPatch.model_validate(patch)
                for identity, columns in state.get("formats", {}).items()
                for column, patch in columns.items()
            }
            for name, state in editor["resources"].items()
        }
        # Revalidate persisted protection metadata at preparation, just like value operations.
        for name, cells in formats.items():
            sheet = package.sheets[name]
            for coordinate in cells:
                index, column = coordinate_to_tuple(coordinate)
                options = (
                    editor["resources"][name]
                    .get("format_options", {})
                    .get(f"r:{index - 1}", {})
                    .get(str(column - 1), {})
                )
                if sheet.structural_readonly(coordinate) or (
                    sheet.protected
                    and coordinate not in sheet.unlocked
                    and not options.get("override_protection")
                ):
                    raise ApplicationError(
                        "EDITOR_CELL_READONLY",
                        "Enable Edit calculations to change formulas or protected cell values.",
                    )
        content = format_book(package.patched(updates), formats)
        if changed_formulas:
            content = invalidate_formula_chain(content)
        patched = WorkbookCells(content)
        if any(
            c.find(NS + "f") is not None for s in patched.sheets.values() for c in s.cells.values()
        ):
            content = recalculate(content)
        return io.BytesIO(content), "browser-edited.xlsx"

    def prepare_metadata(self, output: io.BytesIO) -> dict[str, Any]:
        return {"external_cache": bool(WorkbookCells(output.getvalue()).external)}
