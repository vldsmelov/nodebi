import re
from datetime import date, datetime
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import polars as pl
from defusedxml.common import DefusedXmlException  # type: ignore[import-untyped]
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.datasets import (
    ColumnSchema,
    Dataset,
    DatasetSchema,
    LogicalType,
    logical_to_polars,
)
from app.domain.workbooks import (
    DEFAULT_WORKSHEET_LAYOUT,
    CellValue,
    ColumnInfo,
    WorkbookInfo,
    WorksheetInfo,
    WorksheetLayout,
    WorksheetLoadOptions,
    WorksheetPreview,
)
from app.infrastructure.storage import StorageProvider
from app.infrastructure.worksheet_layout import inspect_structure


class ExcelWorkbookError(Exception):
    """A workbook cannot be inspected safely."""


class WorksheetNotFoundError(ExcelWorkbookError):
    """The requested worksheet is absent."""


def _logical_type(values: list[CellValue]) -> tuple[LogicalType, bool]:
    present = [value for value in values if value is not None]
    nullable = len(present) != len(values)
    if not present:
        return LogicalType.UNKNOWN, True
    value_types = {type(value) for value in present}
    if value_types <= {bool}:
        return LogicalType.BOOLEAN, nullable
    if value_types <= {int}:
        return LogicalType.INTEGER, nullable
    if value_types <= {int, float}:
        return LogicalType.FLOAT, nullable
    if value_types <= {date}:
        return LogicalType.DATE, nullable
    if value_types <= {date, datetime}:
        return LogicalType.DATETIME, nullable
    if value_types <= {str}:
        return LogicalType.STRING, nullable
    return LogicalType.STRING, nullable


def _columns(headers: tuple[Any, ...], data_rows: list[tuple[Any, ...]]) -> list[ColumnInfo]:
    used: set[str] = set()
    labels = [
        (str(value).strip() if value is not None else "") or f"Column {index + 1}"
        for index, value in enumerate(headers)
    ]
    reserved = set(labels)
    result: list[ColumnInfo] = []
    for index, raw_header in enumerate(headers):
        display_name = str(raw_header).strip() if raw_header is not None else ""
        display_name = display_name or f"Column {index + 1}"
        engine_name = display_name
        occurrence = 2
        while engine_name in used:
            engine_name = f"{display_name}__{occurrence}"
            occurrence += 1
            if engine_name in reserved:
                engine_name = display_name
        used.add(engine_name)
        values = [row[index] if index < len(row) else None for row in data_rows]
        logical_type, nullable = _logical_type(values)
        result.append(
            ColumnInfo(
                column_id=f"col_{index + 1}",
                display_name=display_name,
                engine_name=engine_name,
                logical_type=logical_type,
                nullable=nullable,
            )
        )
    return result


class OpenPyxlWorkbookReader:
    def __init__(self, storage: StorageProvider) -> None:
        self.storage = storage

    def _open_workbook(self, storage_key: str):  # type: ignore[no-untyped-def]
        stream = None
        try:
            stream = self.storage.open(storage_key)
            with ZipFile(stream) as archive:
                if (
                    len(archive.infolist()) > 10_000
                    or sum(member.file_size for member in archive.infolist()) > 256 * 1024 * 1024
                ):
                    raise ValueError("Workbook expanded size exceeds the supported limit.")
            stream.seek(0)
            return stream, load_workbook(stream, read_only=True, data_only=True, keep_vba=False)
        except (
            BadZipFile,
            InvalidFileException,
            OSError,
            ValueError,
            KeyError,
            ParseError,
            DefusedXmlException,
        ) as exc:
            if stream is not None:
                stream.close()
            raise ExcelWorkbookError("The uploaded file is not a readable Excel workbook.") from exc

    def inspect(self, storage_key: str) -> WorkbookInfo:
        stream, workbook = self._open_workbook(storage_key)
        try:
            sheets = [
                WorksheetInfo(
                    name=sheet.title,
                    row_count=max((sheet.max_row or 1) - 1, 0),
                    column_count=sheet.max_column or 0,
                )
                for sheet in workbook.worksheets
            ]
            if not sheets:
                raise ExcelWorkbookError("The workbook does not contain any worksheets.")
            return WorkbookInfo(sheets=sheets)
        finally:
            workbook.close()
            stream.close()

    def _sheet_rows(self, stream, sheet, layout: WorksheetLayout):  # type: ignore[no-untyped-def]
        try:
            # Isolated access to the OOXML part path of pinned openpyxl.
            structure = inspect_structure(stream, sheet._worksheet_path)
        except ValueError as exc:
            raise ExcelWorkbookError(str(exc)) from exc
        header_end = layout.header_row + layout.header_rows
        data_start = header_end + layout.skip_data_rows
        data_end = max(header_end - 1, structure.rows - layout.drop_tail_rows)
        missing = [
            cell
            for cell in structure.missing_formula_results
            if layout.header_row <= int(re.sub(r"\D", "", cell)) < header_end
            or data_start <= int(re.sub(r"\D", "", cell)) <= data_end
        ]
        if missing:
            raise ExcelWorkbookError(
                f"Formula has no saved result at {sheet.title}!{missing[0]}. "
                "Recalculate and save the workbook in Excel before uploading."
            )
        width = structure.columns
        if not width:
            return (), iter(()), []
        if width + layout.include_row_number > 1024:
            raise ExcelWorkbookError("Worksheet exceeds the 1024-column limit.")
        if (
            max(0, data_end - data_start + 1) > 250_000
            or max(0, data_end - data_start + 1) * (width + layout.include_row_number) > 2_000_000
        ):
            raise ExcelWorkbookError("Worksheet exceeds the 250000-row or 2000000-cell limit.")
        header_values = list(
            sheet.iter_rows(
                min_row=layout.header_row,
                max_row=header_end - 1,
                max_col=width,
                values_only=True,
            )
        )
        headers = [list(row) for row in header_values]
        for left, top, right, bottom in structure.merges:
            if layout.header_row <= top < header_end:
                value = headers[top - layout.header_row][left - 1] if left <= width else None
                for row in range(top, min(bottom + 1, header_end)):
                    for column in range(left, min(right, width) + 1):
                        headers[row - layout.header_row][column - 1] = value
        levels = [[row[column] for row in headers] for column in range(width)]
        labels = tuple(
            " / ".join(dict.fromkeys(str(value).strip() for value in values if value is not None))
            or None
            for values in levels
        )
        iterator = (
            sheet.iter_rows(
                min_row=data_start,
                max_row=data_end,
                max_col=width,
                values_only=True,
            )
            if data_end >= data_start
            else iter(())
        )
        if layout.include_row_number:
            labels = (*labels, "__source_row__")
            levels.append(["__source_row__"])
            iterator = ((*row, number) for number, row in enumerate(iterator, data_start))
        return labels, iterator, levels

    def preview_sheet(
        self, storage_key: str, sheet_name: str, options: WorksheetLoadOptions
    ) -> WorksheetPreview:
        stream, workbook = self._open_workbook(storage_key)
        try:
            if sheet_name not in workbook.sheetnames:
                raise WorksheetNotFoundError(f'Worksheet "{sheet_name}" does not exist.')
            headers, iterator, levels = self._sheet_rows(
                stream, workbook[sheet_name], options.layout
            )
            sampled = []
            for _, row in zip(range(options.limit + 1), iterator, strict=False):
                sampled.append(row)
            if not headers or (all(value is None for value in headers) and not sampled):
                return WorksheetPreview(sheet_name, [], [], 0, False)
            truncated = len(sampled) > options.limit
            data_rows = sampled[: options.limit]
            columns = _columns(headers, data_rows)
            rows = [
                {
                    column.column_id: _serialize(row[index] if index < len(row) else None)
                    for index, column in enumerate(columns)
                }
                for row in data_rows
            ]
            return WorksheetPreview(sheet_name, columns, rows, len(rows), truncated)
        except (ParseError, DefusedXmlException) as exc:
            raise ExcelWorkbookError("The worksheet contains invalid or unsafe XML.") from exc
        finally:
            workbook.close()
            stream.close()

    def load_sheet(
        self, storage_key: str, sheet_name: str, layout: WorksheetLayout = DEFAULT_WORKSHEET_LAYOUT
    ) -> Dataset:
        stream, workbook = self._open_workbook(storage_key)
        try:
            if sheet_name not in workbook.sheetnames:
                raise WorksheetNotFoundError(f'Worksheet "{sheet_name}" does not exist.')
            headers, iterator, levels = self._sheet_rows(stream, workbook[sheet_name], layout)
            rows: list[tuple[Any, ...]] = []
            for row in iterator:
                if len(rows) >= 250_000 or (len(rows) + 1) * len(headers) > 2_000_000:
                    raise ExcelWorkbookError(
                        "Worksheet exceeds the 250000-row or 2000000-cell limit."
                    )
                rows.append(row)
            if not headers or (all(value is None for value in headers) and not rows):
                return Dataset(pl.DataFrame(), DatasetSchema(()), {"sheet_name": sheet_name})
            inferred = _columns(headers, rows)
            # Polars' row inference samples a prefix. Our logical inference examines
            # all rows and explicitly normalizes mixed values before construction.
            frame = pl.DataFrame(
                [
                    pl.Series(
                        column.engine_name,
                        [_coerce(row[index], column.logical_type) for row in rows],
                        dtype=logical_to_polars(column.logical_type),
                        strict=True,
                    )
                    for index, column in enumerate(inferred)
                ]
            )
            columns = tuple(
                ColumnSchema(
                    column_id=column.column_id,
                    engine_name=column.engine_name,
                    display_name=column.display_name,
                    logical_type=column.logical_type,
                    nullable=frame[column.engine_name].null_count() > 0,
                    metadata={"source_row_number": True}
                    if layout.include_row_number and index == len(inferred) - 1
                    else {"source_position": index, "header_levels": levels[index]},
                )
                for index, column in enumerate(inferred)
            )
            return Dataset(
                frame=frame,
                schema=DatasetSchema(columns),
                metadata={"sheet_name": sheet_name},
            )
        except (ParseError, DefusedXmlException) as exc:
            raise ExcelWorkbookError("The worksheet contains invalid or unsafe XML.") from exc
        finally:
            workbook.close()
            stream.close()


def _serialize(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _coerce(value: Any, logical_type: LogicalType) -> Any:
    if value is None:
        return None
    if logical_type == LogicalType.STRING:
        return str(_serialize(value))
    if logical_type == LogicalType.FLOAT:
        return float(value)
    if logical_type == LogicalType.DATETIME and type(value) is date:
        return datetime.combine(value, datetime.min.time())
    return value
