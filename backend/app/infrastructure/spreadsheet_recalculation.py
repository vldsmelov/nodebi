"""Recalculate a data/formula-only copy; never hand original workbook assets to Calc."""

import copy
import io
import os
import re
import shutil
import signal
import subprocess
import tempfile
from pathlib import Path
from threading import BoundedSemaphore
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import quote_sheetname
from openpyxl.worksheet.filters import AutoFilter

from app.domain.errors import ApplicationError
from app.infrastructure.workbook_cells import NS, WorkbookCells

_gate = BoundedSemaphore(1)
# Deliberately excludes network, DDE, macro/add-in and indirect external evaluation.
SAFE_FUNCTIONS = set(
    [
        "ABS",
        "ACOS",
        "AND",
        "ASIN",
        "ATAN",
        "ATAN2",
        "AVERAGE",
        "AVERAGEIF",
        "AVERAGEIFS",
        "CEILING",
        "CHAR",
        "CHOOSE",
        "CLEAN",
        "CODE",
        "COLUMN",
        "COLUMNS",
        "CONCAT",
        "CONCATENATE",
        "COUNT",
        "COUNTA",
        "COUNTBLANK",
        "COUNTIF",
        "COUNTIFS",
        "DATE",
        "DATEDIF",
        "DATEVALUE",
        "DAY",
        "DAYS",
        "EDATE",
        "EOMONTH",
        "EXACT",
        "EXP",
        "FALSE",
        "FIND",
        "FLOOR",
        "HLOOKUP",
        "HOUR",
        "IF",
        "IFERROR",
        "IFNA",
        "IFS",
        "INDEX",
        "INT",
        "ISBLANK",
        "ISERROR",
        "ISEVEN",
        "ISLOGICAL",
        "ISNA",
        "ISNONTEXT",
        "ISNUMBER",
        "ISODD",
        "ISREF",
        "ISTEXT",
        "LARGE",
        "LEFT",
        "LEN",
        "LN",
        "LOG",
        "LOG10",
        "LOOKUP",
        "LOWER",
        "MATCH",
        "MAX",
        "MAXIFS",
        "MEDIAN",
        "MID",
        "MIN",
        "MINIFS",
        "MINUTE",
        "MOD",
        "MONTH",
        "MROUND",
        "N",
        "NA",
        "NOT",
        "NOW",
        "OR",
        "PI",
        "POWER",
        "PRODUCT",
        "PROPER",
        "QUARTILE",
        "QUARTILE.INC",
        "QUARTILE.EXC",
        "RANK",
        "RANK.EQ",
        "REPLACE",
        "REPT",
        "RIGHT",
        "ROUND",
        "ROUNDDOWN",
        "ROUNDUP",
        "ROW",
        "ROWS",
        "SEARCH",
        "SECOND",
        "SIGN",
        "SMALL",
        "SQRT",
        "STDEV",
        "STDEV.S",
        "STDEV.P",
        "SUBSTITUTE",
        "SUBTOTAL",
        "SUM",
        "SUMIF",
        "SUMIFS",
        "SUMPRODUCT",
        "TEXT",
        "TIME",
        "TIMEVALUE",
        "TODAY",
        "TRIM",
        "TRUE",
        "TRUNC",
        "UPPER",
        "VALUE",
        "VAR",
        "VAR.S",
        "VAR.P",
        "VLOOKUP",
        "WEEKDAY",
        "WEEKNUM",
        "XLOOKUP",
        "XMATCH",
        "XOR",
        "YEAR",
        "YEARFRAC",
    ]
)


def safe_formula(formula: str, external_names: dict[tuple[int, str], str] | None = None) -> str:
    try:
        tokens = Tokenizer(formula if formula.startswith("=") else "=" + formula).items
    except Exception as error:
        raise ApplicationError(
            "EDITOR_FORMULA_UNSUPPORTED",
            "This workbook contains a formula that cannot be recalculated safely.",
        ) from error
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper().removeprefix("_XLFN.")
            if name not in SAFE_FUNCTIONS:
                raise ApplicationError(
                    "EDITOR_FORMULA_UNSUPPORTED",
                    "This workbook contains a formula that cannot be recalculated safely.",
                )
        if (
            token.type == "OPERAND"
            and token.subtype == "RANGE"
            and any(x in token.value for x in ("[", "]", "|", "://"))
        ):
            reference = re.fullmatch(r"'?(?:\[(\d+)\])(.+?)'?!([A-Za-z0-9$:.]+)", token.value)
            key = (int(reference[1]), reference[2].replace("''", "'")) if reference else None
            if not reference or not external_names or key not in external_names:
                raise ApplicationError(
                    "EDITOR_EXTERNAL_FORMULA",
                    "An external reference has no saved values. Your input has not been changed.",
                )
            token.value = f"{quote_sheetname(external_names[key])}!{reference[3]}"
    return "=" + "".join(token.value for token in tokens)


def office_binary() -> str:
    configured = os.environ.get("NODEBI_SPREADSHEET_BINARY")
    binary = configured or shutil.which("libreoffice") or shutil.which("soffice")
    if not binary and os.name == "nt":
        candidate = (
            Path(os.environ.get("PROGRAMFILES", "C:/Program Files"))
            / "LibreOffice/program/soffice.com"
        )
        binary = str(candidate) if candidate.is_file() else None
    if not binary:
        raise ApplicationError(
            "EDITOR_CALCULATION_UNAVAILABLE",
            "The spreadsheet calculation service is unavailable. "
            "Your draft is saved; try again after the service is restored.",
            503,
        )
    return binary


def recalculate(content: bytes) -> bytes:
    if not _gate.acquire(blocking=False):
        raise ApplicationError(
            "EDITOR_CALCULATION_BUSY",
            "Another workbook is being recalculated. Your draft is saved; try again shortly.",
            503,
        )
    try:
        try:
            return _recalculate(content)
        except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
            raise ApplicationError(
                "EDITOR_CALCULATION_UNAVAILABLE",
                "The spreadsheet calculation service is unavailable. Your draft is saved.",
                503,
            ) from error
    finally:
        _gate.release()


def _recalculate(content: bytes) -> bytes:
    binary = office_binary()
    version = subprocess.run(
        [binary, "--version"], capture_output=True, text=True, timeout=15, check=True
    )
    match = re.search(r"(\d+)\.(\d+)\.(\d+)", version.stdout)
    if not match or tuple(map(int, match.groups())) < (26, 2, 5):
        raise ApplicationError(
            "EDITOR_CALCULATION_UNAVAILABLE",
            "The spreadsheet calculation service must be updated before saving this workbook.",
            503,
        )
    package = WorkbookCells(content)
    original = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    calculation = Workbook()
    calculation.remove(calculation.active)  # type: ignore[arg-type]
    calculation.epoch = original.epoch
    calculation.calculation = copy.copy(original.calculation)
    try:
        # Only literal saved cache cells cross this boundary; never paths/connections.
        external_names: dict[tuple[int, str], str] = {}
        for index, (key, values) in enumerate(package.external.items()):
            title = f"_NodeBI_link_{index}"
            while title in package.sheets or title in calculation.sheetnames:
                title += "_"
            target = calculation.create_sheet(title)
            target.sheet_state = "hidden"
            external_names[key] = title
            for coordinate, kind, value in values:
                if not value:
                    continue
                written = target[coordinate]
                if kind in {"n", "b"}:
                    written.value = float(value) if kind == "n" else value == "1"
                else:
                    written.value = value
                    written.data_type = "e" if kind == "e" else "s"
        for sheet in package.sheets.values():
            target = calculation.create_sheet(sheet.name)
            for row_number in sheet.hidden_rows:
                target.row_dimensions[row_number].hidden = True
            if sheet.auto_filter is not None:
                target.auto_filter = AutoFilter.from_tree(sheet.auto_filter)
            incoming = original[sheet.name]
            cell_value: Any
            for cells in incoming.iter_rows(max_row=sheet.rows, max_col=sheet.columns):
                for cell in cells:
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            raise ApplicationError(
                                "EDITOR_FORMULA_UNSUPPORTED",
                                "Array or data-table formulas cannot be recalculated "
                                "in this editor.",
                            )
                        cell_value = safe_formula(cell.value, external_names)
                    else:
                        cell_value = cell.value
                    written = target.cell(cell.row, cell.column, cell_value)
                    written.data_type = cell.data_type
                    written.number_format = cell.number_format
            for name, definition in incoming.defined_names.items():
                copied = copy.copy(definition)
                copied.attr_text = safe_formula(definition.attr_text or "", external_names)[1:]
                target.defined_names[name] = copied
        for name, definition in original.defined_names.items():
            copied = copy.copy(definition)
            copied.attr_text = safe_formula(definition.attr_text or "", external_names)[1:]
            calculation.defined_names[name] = copied
        with tempfile.TemporaryDirectory(prefix="nodebi-calc-") as folder:
            root = Path(folder)
            profile = root / "profile"
            (profile / "user").mkdir(parents=True)
            (profile / "user/registrymodifications.xcu").write_text(
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
                '<item oor:path="/org.openoffice.Office.Calc/Content/Update">'
                '<prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item>'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
                "</oor:items>",
                encoding="utf-8",
            )
            source = root / "input.xlsx"
            calculation.save(source)
            out = root / "output"
            out.mkdir()
            process = subprocess.Popen(
                [
                    binary,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--norestore",
                    "--nodefault",
                    "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML",
                    "--outdir",
                    str(out),
                    str(source),
                ],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                start_new_session=os.name != "nt",
            )
            try:
                process.communicate(timeout=90)
            except subprocess.TimeoutExpired as error:
                if os.name == "nt":
                    subprocess.run(
                        ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                        capture_output=True,
                        check=False,
                    )
                else:
                    os.killpg(process.pid, signal.SIGKILL)  # type: ignore[attr-defined]
                process.communicate()
                raise ApplicationError(
                    "EDITOR_CALCULATION_TIMEOUT",
                    "Recalculation took too long. Your browser draft is saved; try again.",
                    503,
                ) from error
            result = out / "input.xlsx"
            if process.returncode != 0 or not result.is_file():
                raise ApplicationError(
                    "EDITOR_CALCULATION_FAILED",
                    "The workbook could not be recalculated. Your input has not been changed.",
                )
            calculated = WorkbookCells(result.read_bytes())
            updates: dict[str, dict[str, tuple[str, str, bool]]] = {}
            for sheet in package.sheets.values():
                computed = calculated.sheets.get(sheet.name)
                if computed is None:
                    raise ApplicationError(
                        "EDITOR_CALCULATION_FAILED",
                        "The workbook could not be recalculated. Your input has not been changed.",
                    )
                for coordinate, cell in sheet.cells.items():
                    if cell.find(NS + "f") is None:
                        continue
                    computed_cell = computed.cells.get(coordinate)
                    if computed_cell is None or computed_cell.find(NS + "v") is None:
                        raise ApplicationError(
                            "EDITOR_CALCULATION_FAILED",
                            "A formula has no calculated result. Your input has not been changed.",
                        )
                    kind = computed_cell.get("t", "n")
                    if kind == "e" and (
                        cell.get("t") != "e"
                        or package.value(cell) != calculated.value(computed_cell)
                    ):
                        raise ApplicationError(
                            "EDITOR_FORMULA_ERROR",
                            f"Recalculation produced an error at {sheet.name}!{coordinate}. "
                            "Correct the input and save again.",
                        )
                    text = computed_cell.findtext(NS + "v", "")
                    if kind in {"s", "inlineStr"}:
                        kind, text = "str", calculated.value(computed_cell)
                    updates.setdefault(sheet.path, {})[coordinate] = (kind, text, True)
            return package.patched(updates)
    finally:
        original.close()
        calculation.close()
